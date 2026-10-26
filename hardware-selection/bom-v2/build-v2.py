"""Build the September 8 sourcing snapshot. Only writes inside bom-v2.

Prices are observations/allowances, not checkout quotes. V1 stays historical.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import Counter
import copy
import json
import csv

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DATE = '2026-09-08'
rows = []

def money(value):
    return float(Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))

def add(id, group, name, model, used, buy, unit, price, basis, state, channel,
        url='', code='', moq=None, stock=None, notes='', dimensions='', import_ok=False):
    rows.append(dict(id=id, category=group, item=name, model=model,
        installed_quantity=used, purchase_quantity=buy, unit=unit,
        unit_price_or_allowance_cny=price, price_basis=basis,
        line_cash_cny=money(Decimal(str(price))*Decimal(str(buy))),
        selection_status=state, channel=channel, url=url, lcsc_code=code,
        observed_moq=moq, observed_stock=stock, notes=notes,
        mechanical_constraints=dimensions, core_import_eligible=import_ok,
        purchase_confirmed=False))

LIVE='本轮网页价'; SHOT='历史用户截图价'; BUDGET='预算占位'; INCLUDED='已含在其他行'
LC='立创商城'; MRO='立创工业品'; DFR='DFRobot 官网'; TBD='待找小批量渠道'

add('E01','显示与控制','2.4英寸 TFT＋EC11 屏幕','TOUGLESY ST7789 320×240 SPI',1,1,'块',32.40,SHOT,'用户已确认；待实测','既定淘宝商品',
    notes='独立屏幕，不带触摸；原商品截图已归档，未核结算价。',dimensions='PCB 标称91×44mm；孔距84.5×37.5mm、Ø3。有效显示区与控件高度待测。')
add('E02','执行机构','金属齿轮180度定位舵机及随包舵盘螺丝','MG90S（用户选中店铺款）',8,8,'个',10,SHOT,'用户已确认；待单腿负载测试','既定淘宝商品',
    notes='每腿2个；试样使用其中2个，整机仍只计8个。不能将原厂资料视作此店铺实物保证。',dimensions='黑色实物款；安装耳、轴中心、舵盘孔位待实测，复用随包舵盘。')
add('E03','显示与控制','ESP32-S3 开发板','LCKFB-ESP32S3R8N8',1,1,'块',35.90,LIVE,'用户已确认',LC,
    'https://item.szlcsc.com/22034693.html','C20626143',1,1488,
    '8MB Flash＋8MB PSRAM；附未焊1×20P排针，需焊接。按普通价，未套用学生价30.515元。13.3g为商品毛重。',
    '板体、USB口位置、天线禁布及GPIO分配待原理图/尺寸图核对；不沿用旧N16R8占位。',True)
add('E04','显示与控制','16路 PWM 舵机信号模块','PCA9685 成品模块，具体SKU待选',1,1,'块',12,BUDGET,'待选现货商品',TBD,
    notes='要成品模块、逻辑VCC=3.3V、V+独立。仅提供PWM；8舵机主电流从独立配电板走。12元不是查到的报价。',
    dimensions='模块板框、安装孔和OE接口待定；不能把裸芯片尺寸用作模块尺寸。')
add('E05','语音','I²S 数字麦克风模块','SEN0327',1,1,'块',15,LIVE,'本轮推荐候选',DFR,
    'https://www.dfrobot.com.cn/goods-2567.html',stock='有货',
    notes='1.8–3.3V供电，接3.3V；24bit I²S，底部拾音。页面未注明INMP441，不强行认定芯片型号。官网品牌栏为其它品牌。',
    dimensions='PCB Ø14mm；厚度、声孔位置与排针高度待核；导音孔对准底部。')
add('E06','语音','I²S 单声道功放模块','DFR0954 / MAX98357A',1,1,'块',30,LIVE,'本轮推荐候选',DFR,
    'https://www.dfrobot.com.cn/goods-3573.html',stock='有货',
    notes='3.3–5V供电；本方案接逻辑5V。PH2.0接口或焊盘；与所选喇叭1.25mm插头不直配，改焊线或核对转接件。BTL的OUT+/OUT−均不能接地。',
    dimensions='官网文本尺寸栏为空，板框/高件/插头空间待补。')
add('E07','语音','带线小喇叭 4Ω 3W','24MM-4Ω3W-2P1.25',1,1,'个',9.69,LIVE,'本轮推荐候选',LC,
    'https://item.szlcsc.com/53008620.html','C50387237',1,89,
    '额定3W；2pin 1.25mm线端插头，不能仅靠间距认定连接器系列。与DFR0954采用焊线连接方案，实际音量/背腔待测。',
    '参数表Ø23.6×10.4mm，简介约24×10.5mm；替换V1中Ø28×7mm占位。',True)
add('P01','电源','定制成品2S动力电池包，含保护均衡和XT30配对接头','2S1P 7.4V 1000mAh；持续≥10A为定制目标',1,1,'包',80,BUDGET,'用户确认定制规格；供货和报价未确认','既定定制入口',
    'https://item.taobao.com/item.htm?id=966262592635',
    notes='每节保护、已连接被动均衡、共口P+/P−，允许充电≥1A为询价要求。80元未获商家报价；不能套用14.70元普通对讲机包价格。',
    dimensions='≤75×40×20mm、≤80g均是目标，含保护板和包膜；成品尺寸/重量未知。')
add('P02','电源','2S 8.4V Type-C 充电模块','TELESKY MK021163 TYPE-C接口款',1,1,'块',14.13,SHOT,'用户已确认款式；充电电流待核','既定天猫商品',
    'https://detail.tmall.com/item.htm?id=709098048793&skuId=5148568746386',
    notes='5V2A是输入条件；充电电流目标0.4–0.5A，实际CC值未知。首版关机充电。',dimensions='标称28×17mm；厚度、孔位及USB口坐标待测。')
for id, label, status in [('P03','舵机路独立降压模块','推荐候选；6A目标待实测'),('P04','逻辑/音频路独立降压模块','同款备选计价；仍可寻找更便宜2A级模块')]:
    add(id,'电源',label,'DFR0753',1,1,'块',55,LIVE,status,DFR,
        'https://www.dfrobot.com.cn/goods-2971.html',stock='有货',
        notes='6–14V输入、5V±1%输出、最大8A为厂家标称；未证明在低电量2S及壳内连续达到该电流。两块独立供电、共地，5V输出不并联。低于6V不在输入规格内。',
        dimensions='板体和焊盘尺寸待图纸；散热、导线焊接与固定待设计。')
add('P05','电源','USB-A 转 Type-C 数据/充电线','5V≥2A，建议USB2.0数据线，约1m',1,1,'条',5,BUDGET,'待选现货；已有合格线可复用',TBD,
    notes='绿联60116/C7462745已核对规格但立创缺货，不混入可购清单。开发和充电可以分时使用1条，不支持同时连接即需另加。')
add('P06','电源','电源操作开关','优先小电流控制负载开关，形式待设计',1,1,'个',5,BUDGET,'待原理图与面板形式',TBD,
    notes='与F05负载开关配套；若改为直接断开动力电源，需明确DC额定及接线。不能用AC额定值推断DC大电流能力。')
add('F01','保护与滤波','MINI 汽车插片保险丝 7.5A 32VDC','Littelfuse 029707.5WXNV',1,5,'个',0.6162,LIVE,'PCB方案候选；熔断值未冻结',LC,
    'https://item.szlcsc.com/190335.html','C178942',5,5935,
    '装机1个，计划至少1个备件；MOQ5。7.5A仅试制起点，需对照线材与时间电流曲线/舵机浪涌，不是7.5A立即跳闸。')
add('F02','保护与滤波','MINI 保险丝 PCB 座','Littelfuse 01530008Z',1,1,'个',24.35,LIVE,'PCB方案候选；需设计载板',LC,
    'https://item.szlcsc.com/208026.html','C206907',1,534,
    '32V/20A PCB型，与MINI保险丝配套；不是带线串联座。需近电池的载板布局与合格铜箔电流路径。',
    '按规格书焊盘/保险丝拔出高度设计；直线束方案应换有明确DC额定的带线座并扣除此行。')
add('F03','保护与滤波','舵机5V路储能电容 1000µF 10V','KNSCHA SHW1000UF10V8X12',2,5,'个',0.531608,LIVE,'本轮推荐；接5V支路',LC,
    'https://item.szlcsc.com/505124.html','C494548',5,7975,
    '低阻直插；2个为初始配置，按压降/稳定性实测调整。不要据此替代降压模块要求的本地电容。',
    'Ø8×12mm，脚距3.5mm；保留极性和装配间隙。',True)
add('F04','保护与滤波','逻辑5V路储能电容 470µF 16V','JIERR ERB16V470M8X12',1,10,'个',0.38855,LIVE,'本轮推荐；接逻辑5V',LC,
    'https://item.szlcsc.com/54681555.html','C51953430',10,480,
    '采购MOQ10，装机暂1。最终位置和数量按原理图及降压稳定性核对。', 'Ø8×12mm。',True)
add('F05','保护与滤波','高边负载开关、栅极控制及必要小件','MOSFET/控制器待原理图选型',1,1,'套',5,BUDGET,'待电路设计',TBD,
    notes='总开关应断开两路工作负载，充电支路保留；低电量停止动作后的保持/断电策略待设计。')
add('F06','保护与滤波','电池电压采样及RC滤波','分压电阻/RC，阻值与GPIO待定',1,1,'套',2,BUDGET,'待电路设计',TBD,
    notes='满电8.4V也须低于所用ADC输入范围；低电量策略必须计入负载压降、回差和每节电池保护。不是只依赖软件监测。')
add('F07','保护与滤波','本地去耦、PWM默认关闭及接口小件','电容/电阻/必要电平处理待原理图',1,1,'套',3,BUDGET,'待电路设计',TBD,
    notes='含PCA9685 OE上电默认禁用等必要器件预算；与F03/F04大电容分开，PCB裸器件尚未形成正式BOM。')
add('W01','接线与连接器','开发板插座 1×20P 2.54mm 直排母','JILN22850120ANG1SYA01',2,5,'条',1.0612,LIVE,'推荐用于可拔插开发板',LC,
    'https://item.szlcsc.com/426399.html','C429947',5,2280,
    '开发板附公排针，不重复采购公针；该排母单针额定3A，不用作8舵机公共电流通路。',
    '塑高8.5mm；还要加焊接、板厚与开发板元件高度。',True)
add('W02','接线与连接器','10P 软硅胶信号连接线套装','LCKFB-JUMP-WIRES-10P',3,3,'套',4.90,LIVE,'候选；公母头和长度待核',LC,
    'https://item.szlcsc.com/56016470.html','C53063971',1,348,
    '每套10根，24AWG。共30根是接线预留，未形成线束表；文本未核实公母及长度，暂不进入导入表。不可作电池/舵机主电源线。')
for id, color in [('W03','红'),('W04','黑')]:
    add(id,'接线与连接器',f'主干{color}色硅胶线','18AWG 多股，按米零售',1,1,'米',5,BUDGET,'待选小包装商品',TBD,
        notes='采购1m供裁切和返工；按10A目标、实际长度、线束成束温升/压降验证。电池自带短引线仍由厂家核定。')
for id, color in [('W05','红'),('W06','黑')]:
    add(id,'接线与连接器',f'分支{color}色硅胶线','20AWG 多股，按米零售',1,1,'米',4,BUDGET,'待选小包装商品',TBD,
        notes='采购1m供短分支；最终长度随布局，单段电流/接点额定分别核算，不用杜邦线承载舵机总电流。')
add('W07','接线与连接器','舵机3线端口及模块排针/转接','2.54mm 单排3P×8；另含模块所需短排针',1,1,'套',4,BUDGET,'待线序和载板设计',TBD,
    notes='复用MG90S自带线；逐个端口核对GND/5V/PWM及接点额定。全部舵机电源由载板粗电流通路分配。功放喇叭先按焊线，未额外算PH2.0转接线。')
add('W08','接线与连接器','XT30 主电源配对接头','随P01定制包包含配对两半',1,0,'对',0,INCLUDED,'随电池要求；无额外采购','电池供应商',
    notes='不要重复买2对。AMASS XT30U-M/C99101与XT30U-F/C99102已找到立创MRO入口，但各MOQ5且首档价未核，额外备件待需时再计。')
add('W09','接线与连接器','绝缘热缩管小包装','约2/3/5/8mm混合，尺寸依焊点',1,1,'包',3,BUDGET,'待小包装商品',TBD,
    notes='未选200m整卷；覆盖电源接头、喇叭改焊点，直径按焊后包络核对。')
add('W10','接线与连接器','原型信号转接洞洞板','约4×6cm FR4 双面洞洞板',1,1,'块',3,BUDGET,'试制辅料；待商品',TBD,
    notes='用于信号接线试制；不能默认普通洞洞板焊盘满足8舵机动力电流。正式载板另计P09。')
add('P09','PCB与输入附件','定制载板/配电PCB首次小批打样','目标5片小批，单机用1片',1,1,'批',30,BUDGET,'待原理图、板框与布局','嘉立创 PCB',
    notes='计整批现金，不除以5；未含SMT或生产文件。F05/F06/F07只是器件预留。')
add('P10','PCB与输入附件','成品 USB-A 电源适配器','固定5V输出≥2A，有合规标识',1,1,'个',20,BUDGET,'待现货；已有合格充电头可复用',TBD,
    notes='仅作Type-C充电板输入。已查20328/CD104立创缺货且文字参数不足，未定为兼容型号；不用裸露市电电源替代。')
add('M01','结构辅料','M2×6 圆头十字机牙螺丝','顺天 PM2X6 镀镍',20,100,'个',0.010735,LIVE,'样件候选；长度/数量未冻结',MRO,
    'https://item.szlcsc.com/mro/332656.html','C357534',100,7300,
    'MOQ/增量均100，装机20只是起始清单；不能当舵盘中心自攻螺丝。')
add('M02','结构辅料','M3×8 圆头十字机牙螺丝','顺天 PM3X8 镀镍',12,100,'个',0.013775,LIVE,'样件候选；长度/数量未冻结',MRO,
    'https://item.szlcsc.com/mro/332666.html','C357544',100,19600,
    'MOQ/增量均100。屏幕/外壳具体使用M2还是M3，以孔位和支架设计为准。')
add('M03','结构辅料','M2 标准六角螺母','约20个，捕获螺母槽方案',1,1,'包',2,BUDGET,'待商品与结构',TBD,
    notes='包内目标20个；外形按实际标准核槽，不拿SMT焊接螺母替代。')
add('M04','结构辅料','M3 铜热熔嵌件或捕获螺母','约12个，固定方式二选一',1,1,'包',3,BUDGET,'待打印材料与结构',TBD,
    notes='外径/长度未冻结。不能把检索到的SMT贴片铜螺母当热熔螺母；两种方案只计一份。')
add('M05','结构辅料','M3 平垫片','约12个',1,1,'包',1,BUDGET,'待商品与结构',TBD)
add('M06','结构辅料','屏幕/模块隔离垫柱','约4个，高度与孔径待测',1,1,'套',2,BUDGET,'待高度链；可并入打印件',TBD,
    notes='若直接打印成支架，删除这笔外购预算，勿重复计。')
add('M07','结构辅料','软脚垫','约12mm硅胶/橡胶，或TPU打印',4,4,'个',2,BUDGET,'待脚底尺寸与工艺',TBD,
    notes='整机4只；若改TPU打印，取消本行并更新打印报价。')
add('M08','结构辅料','电池魔术贴束带','约10–12mm宽、150mm长',1,1,'条',2,BUDGET,'待电池实际尺寸',TBD)
add('M09','结构辅料','缓冲及导音泡棉','约100×100×1mm，材料按用途区分',1,1,'片',2,BUDGET,'待商品与布局',TBD,
    notes='电池缓冲与麦克风导音按实际用途选材，不能堵住拾音孔。')
add('M10','结构辅料','小扎带','约10根',1,1,'包',1,BUDGET,'待小包装商品',TBD)
add('M11','打印','整套外壳、腿与安装件打印','奶油白/薄荷绿；原色FDM基础处理',1,1,'套',90,BUDGET,'未生成生产文件','待打印服务报价',
    notes='不按广告每克价推算；精细喷漆另询。结构清单沿用V1计划，具体装配待深化。')
add('M12','打印','屏幕样框、单腿与公差试样','第一轮试配',1,1,'轮',20,BUDGET,'待实物尺寸后建模','待打印服务报价',
    notes='额外样片费用，不与整套打印重复。使用8个舵机中的2个。')
add('L01','费用预留','多供应商运费','最终按各订单结算',1,1,'项',25,BUDGET,'预留','多渠道')
add('L02','费用预留','补件与试制返工','未发生支出',1,1,'项',40,BUDGET,'预留','项目预留',
    notes='不含购置工具、打印机、长期联网语音服务和额外复杂返工。')

alternatives = [
    dict(model='DFR0954',code='C21292518',url='https://item.szlcsc.com/22712835.html',price_cny=37.75,stock=0,
         decision='立创缺货；同款官网本轮30元有货，主清单采用官网价。'),
    dict(model='Servo Driver HAT',code='C5374673',url='https://item.szlcsc.com/6230032.html',price_cny=100.93,stock=0,
         decision='立创缺货且较贵。官方65×30mm、PCA9685、3.3V逻辑，有预留I2C接口可接其他主控；板载5V仅3A，不足作八舵机电源。',
         official_url='https://www.waveshare.net/shop/Servo-Driver-HAT.htm'),
    dict(model='UGREEN 60116',code='C7462745',url='https://item.szlcsc.com/mro/8495052.html',price_cny=16.75,stock=0,
         decision='USB-A转C，USB2.0数据，1m/3A；规格已核，立创缺货，未纳入现货导入。'),
    dict(model='UGREEN 20328 / CD104',code='C53896762',url='https://item.szlcsc.com/mro/57092397.html',price_cny=38.65,stock=0,
         decision='搜索价/库存；详情参数完善中且通用描述不足以确认此SKU各口5V电流，暂不推荐。'),
    dict(model='AMASS XT30U-M',code='C99101',url='https://item.szlcsc.com/mro/100314.html',price_cny=None,stock=None,moq=5,
         decision='标准XT30线端公半边；首档价未核，起价1.0251不能当MOQ价。电池要求自带配对件，默认额外采购0。'),
    dict(model='AMASS XT30U-F',code='C99102',url='https://item.szlcsc.com/mro/100315.html',price_cny=None,stock=None,moq=5,
         decision='标准XT30线端母半边；首档价未核，起价0.9241不能当MOQ价。核对配对、极性、线径后再增备件。'),
    dict(model='DFR1202',code='',url='https://www.dfrobot.com.cn/goods-4071.html',price_cny=68,stock='有货',
         decision='6–24V→5V最大16A，30×26×8mm；高负载依赖背面散热焊盘/载板铜箔。DFR0753温升不足时再评估，不与它重复计价。'),
    dict(model='JDT 01530008.J',code='C55141810',url='https://item.szlcsc.com/58877668.html',price_cny=4.64,stock=372,
         decision='低价MINI PCB保险座候选，但完整DC额定与图纸未核；不能继承Littelfuse同形产品的32V/20A参数。'),
]

rejections = [
    ('K7805-2000R3 / K12T-6A-P / XL4016','输入最低约8V、8.3V、8V等；不能覆盖2S电池放电范围。',''),
    ('TPS5450-LZM5450-5V / C19673926','页面建议输入比输出高≥2V，免散热建议<15W；不能据标题5A认作6V输入的5V6A方案。','https://item.szlcsc.com/mro/20980610.html'),
    ('DFR1015','7.5–30V输入、需压差；排除用于覆盖本项目整个2S运行区间。','https://www.dfrobot.com.cn/goods-3551.html'),
    ('DFR0756','虽写6–32V输入，但要求≥2V压差；并非6V→稳定5V的保证。','https://www.dfrobot.com.cn/goods-2970.html'),
    ('DFR0571 / DFR0570 / DFR0569','分别最低8V输入、固定3.3V输出、5V仅600mA，不满足本轮逻辑5V2A目标。',''),
    ('MAX98357AETE+T / C910544','裸芯片，不能作为带接线接口的MAX98357A成品功放模块采购。',''),
    ('INMP441裸芯片 / EV_INMP441预售板','裸芯片不是麦克风成品板；预售板资料不全，未选。',''),
    ('Motor Driver HAT / C359961','直流电机驱动HAT不能替代本项目的多路舵机PWM模块。',''),
    ('Hobbywing UBEC 10A-Car','查到款式输出6/7.4/8.4V，没有本方案要求的5V档。',''),
    ('300m电线卷 / 200m热缩管 / SMT铜螺母','前两项不适合单机采购；SMT螺母不可混作3D打印热熔嵌件。',''),
]
old = json.loads((ROOT/'hardware-selection/bom-v1/bom-data.json').read_text())
dims = copy.deepcopy(old['dimensions'])
dim_updates = {
    'D07': ['LCKFB-ESP32S3R8N8 主控板','已选N8R8；附未焊公排针；排母塑高8.5mm','板体长宽高、USB口、排针列距、天线禁布、GPIO保留','托板、USB开孔、天线空间','立创C20626143及官方开源入口','https://wiki.lckfb.com/zh-hans/esp32s3r8n8/open-source/','板型已确认；具体引脚/机械尺寸待核'],
    'D08': ['喇叭 24MM-4Ω3W-2P1.25','参数表Ø23.6×10.4mm；简介约24×10.5mm','实物压边、引线出口、背腔；1.25mm端子改焊或转接','声孔、支架、背腔；旧Ø28×7占位作废','立创C50387237','https://item.szlcsc.com/53008620.html','推荐候选，未更新CAD'],
    'D09': ['SEN0327 I²S 麦克风','PCB Ø14mm，底部拾音','厚度、拾音孔坐标、接线与导音通道','收音孔、泡棉和麦克风座','DFRobot商品页','https://www.dfrobot.com.cn/goods-2567.html','推荐候选，未更新CAD'],
    'D10': ['PCA9685 / DFR0954 / 两路DFR0753','PCA成品SKU未定；功放推荐；降压一块推荐一块备选','板框、孔位、焊盘、高件、散热与线束','可更换托板、配电PCB与通风','本轮选型报告','README.md','不能把芯片封装尺寸当成品板尺寸'],
}
for row in dims:
    if row[0] in dim_updates:
        row[1:] = dim_updates[row[0]]

totals = {
    'scenario': '两路各用1块DFR0753的备选计价方案；未定物料用占位预算；不是已批准采购单',
    'reference_cash_cny': money(sum(Decimal(str(r['line_cash_cny'])) for r in rows)),
    'web_observed_items_cny': money(sum(Decimal(str(r['line_cash_cny'])) for r in rows if r['price_basis']==LIVE)),
    'historical_screenshot_items_cny': money(sum(Decimal(str(r['line_cash_cny'])) for r in rows if r['price_basis']==SHOT)),
    'allowances_cny': money(sum(Decimal(str(r['line_cash_cny'])) for r in rows if r['price_basis']==BUDGET)),
    'core_import_subset_cny': money(sum(Decimal(str(r['line_cash_cny'])) for r in rows if r['core_import_eligible'])),
    'previous_reference_cash_cny': 619.53,
}
totals['change_vs_v1_cny'] = money(Decimal(str(totals['reference_cash_cny']))-Decimal('619.53'))

docs = [
    ('LCKFB开发板官方开源资料','https://wiki.lckfb.com/zh-hans/esp32s3r8n8/open-source/','入口已查到；本轮未完成原理图/尺寸图引脚核对。'),
    ('LCKFB开发板规格书','https://atta.szlcsc.com/upload/public/pdf/source/20241206/8BAF30E9665AF2DA28B66306BACEE9A0.pdf','商品页链接，PDF内容未作为本轮已核实GPIO尺寸依据。'),
    ('SEN0327资料链接','https://www.dfrobot.com.cn/images/upload/File/20191224164325rhavtz.pdf','商品页给出的资料链接。'),
    ('DFR0954 Wiki','https://wiki.dfrobot.com.cn/_SKU_DFR0954_I2S_Amplifier_Module','产品资料入口。'),
    ('喇叭规格书','https://atta.szlcsc.com/upload/public/pdf/source/20250826/91ADA97737CF438E235DF2C37B9C3DFA.pdf','商品页链接；本轮尺寸来自商品参数表。'),
    ('PCA9685成品模块教程','https://wiki.lckfb.com/zh-hans/coloreasyduino/module/control/16-ch-servo-drive-module.html','模块形式参考，不是现货采购SKU。'),
]
data = dict(schema_version=2,date=DATE,currency='CNY',scope='单机样机全套模块、接线、结构辅料、PCB/打印和试制预算',
    user_correction='已确认的是舵机而非舵机电源；只有电池包规格和充电板款式确认，两路降压未确认。',
    price_policy='每行按建议采购数量（含MOQ）乘单价后四舍五入到分，再相加；网页价不含最终运费及优惠。搜索起价不视为首档价。',
    inventory_policy='库存和网页价为本轮读取时的记录，导入/结算前需再次核对。未下单，未改购物车。',
    totals=totals,bom=rows,alternatives=alternatives,rejected_candidates=[dict(model=x,reason=y,url=z) for x,y,z in rejections],
    dimensions=dims,planned_print_parts=old['planned_print_parts'],documentation=[dict(name=x,url=y,scope=z) for x,y,z in docs],
    previous_bom='hardware-selection/bom-v1/README.md',
    workbook='hardware-selection/bom-v2/桌面助手-全套配件选型V2.xlsx',
    core_import_workbook='hardware-selection/bom-v2/立创BOM导入-已核料号子集.xlsx',
    core_import_limit='仅5种明确料号的元器件商城条目；不包含缺货件、MRO、未核线材、需载板决定的保险座/保险及未定成品模块。未上传验证。')
(HERE/'bom-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

headers=['编号','分类','部件','型号/规格','单机用量','建议采购量','单位','单价/预留(元)','本行现金(元)','价格依据','选型状态','渠道','立创编号','已查MOQ','已查库存','商品链接','接口/采购说明','尺寸/结构影响']
def flat(r):
    return [r[k] for k in ['id','category','item','model','installed_quantity','purchase_quantity','unit','unit_price_or_allowance_cny','line_cash_cny','price_basis','selection_status','channel','lcsc_code','observed_moq','observed_stock','url','notes','mechanical_constraints']]
wb = Workbook()
overview = wb.active
overview.title='阅读说明'
overview.append(['项目','说明 / 金额'])
for line in [
    ('日期',DATE),('已确认','屏幕、8个MG90S、定制电池规格、TELESKY充电模块、LCKFB-ESP32S3R8N8'),
    ('本轮成果','立创及MRO优先查货；麦克风/功放/降压有具体官网候选，其余缺口见全套BOM。'),
    ('方案口径',totals['scenario']),('整机参考现金',totals['reference_cash_cny']),
    ('其中：网页观察价',totals['web_observed_items_cny']),('其中：历史截图价',totals['historical_screenshot_items_cny']),
    ('其中：占位预算',totals['allowances_cny']),('立创导入子集金额',totals['core_import_subset_cny']),
    ('相比V1参考现金变化',totals['change_vs_v1_cny']),('采购状态','没有下单或上传BOM；全部均非结算报价。'),
    ('价格口径',data['price_policy']),('接口重点','SEN0327接3.3V；功放喇叭端为BTL；1.25mm喇叭插头与PH2.0不直配。'),
    ('供电重点','2S经两路降压；首版关机充电。DFR0753最低6V输入，低电量负载压降与截止策略待验证。'),
    ('尚未完成','PCA成品低价现货SKU、廉价逻辑降压、总开关和采样器件、线束小包装及最终紧固规格。'),
    ('数据更新','金额为已复核的静态快照。修改数量/单价后请重算，或修改build-v2.py后重新生成；不会自动联动。'),
    ('工具/服务边界','未含购置焊接工具、打印机、长期语音服务或新增传感器；PCB器件级BOM要在原理图完成后生成。'),
]: overview.append(line)

def sheet(name, headings, content):
    ws=wb.create_sheet(name); ws.append(headings)
    for row in content: ws.append(row)
    return ws

allws=sheet('全套BOM',headers,[flat(r) for r in rows])
allws.append(['','合计（含预算）','','','','','','',totals['reference_cash_cny']])
lcws=sheet('立创与MRO有货候选',headers,[flat(r) for r in rows if r['channel'] in [LC,MRO]])
extws=sheet('外部官网候选',headers,[flat(r) for r in rows if r['channel']==DFR])
pending=sheet('待定与预算',headers,[flat(r) for r in rows if r['price_basis']==BUDGET])
sheet('备选与缺货',['型号','立创编号','参考价(元)','查到库存','链接','本轮结论'],
    [[a['model'],a['code'],a['price_cny'],a['stock'],a['url'],a['decision']] for a in alternatives])
sheet('排除记录',['型号','排除理由','来源链接'],rejections)
sheet('尺寸待办',['编号','部件','已有依据','待核内容','结构影响','依据来源','链接/资料','状态'],dims)
sheet('打印分件计划',['编号','部件','预计数量','配色/材料方向','工程约束','状态'],old['planned_print_parts'])
sheet('资料入口',['名称','链接','核验范围'],docs)

def style(book):
    for ws in book.worksheets:
        ws.freeze_panes='D2' if ws.max_column>10 else 'A2'
        ws.sheet_view.showGridLines=False
        ws.auto_filter.ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        ws.row_dimensions[1].height=30
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='365A50'); cell.font=Font(name='PingFang SC',bold=True,color='FFFFFF',size=11)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
        for rr in ws.iter_rows(min_row=2):
            for c in rr:
                c.font=Font(name='PingFang SC',size=10,color='24342F')
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if c.row%2==0: c.fill=PatternFill('solid',fgColor='EFF5F1')
                if isinstance(c.value,str) and c.value.startswith('https://'):
                    c.hyperlink=c.value; c.font=Font(name='PingFang SC',size=10,color='176B87',underline='single')
            ws.row_dimensions[rr[0].row].height=78 if ws.max_column>10 else 60
        if ws.max_column==18:
            widths=[9,17,28,36,12,13,8,15,15,20,34,20,17,12,14,46,70,58]
            for col,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(col)].width=width
            for rr in ws.iter_rows(min_row=2):
                rr[7].number_format='0.000000'; rr[8].number_format='0.00'
        else:
            for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=45 if col>1 else 24
        ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_title_rows='1:1'
    if '阅读说明' in book.sheetnames:
        book['阅读说明'].column_dimensions['A'].width=30
        book['阅读说明'].column_dimensions['B'].width=115
        for row in book['阅读说明'].iter_rows(min_row=2):
            if isinstance(row[1].value,(int,float)): row[1].number_format='0.00'

style(wb)
wb.save(HERE/'桌面助手-全套配件选型V2.xlsx')

core=[r for r in rows if r['core_import_eligible']]
iw=Workbook(); iw.active.title='BOM导入'
iw.active.append(['客户物料编号','商品编号','制造商型号','需求数量','备注'])
for r in core:
    iw.active.append([r['id'],r['lcsc_code'],r['model'],r['purchase_quantity'],r['item']+'；'+r['notes']])
style(iw); iw.active.column_dimensions['E'].width=90
iw.save(HERE/'立创BOM导入-已核料号子集.xlsx')
with (HERE/'立创BOM导入-已核料号子集.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.writer(f); writer.writerow([c.value for c in iw.active[1]])
    writer.writerows([[c.value for c in row] for row in iw.active.iter_rows(min_row=2)])

# Reopen generated artifacts and verify quantities, decimals, IDs and source scope.
assert len({r['id'] for r in rows}) == len(rows)
assert len(core)==5 and all(r['observed_stock'] and r['channel']==LC for r in core)
assert all(r['purchase_quantity']>=r['installed_quantity'] for r in rows if r['id']!='W08')
assert all(not r['lcsc_code'] or r['lcsc_code'].startswith('C') for r in rows)
for r in rows:
    if r['observed_moq']:
        assert r['purchase_quantity']>=r['observed_moq']
    assert r['line_cash_cny']==money(Decimal(str(r['unit_price_or_allowance_cny']))*Decimal(str(r['purchase_quantity'])))
assert money(sum(Decimal(str(totals[k])) for k in ['web_observed_items_cny','historical_screenshot_items_cny','allowances_cny']))==totals['reference_cash_cny']
check=load_workbook(HERE/'桌面助手-全套配件选型V2.xlsx',data_only=True)
assert check['全套BOM'].max_row == len(rows)+2
assert money(sum(Decimal(str(check['全套BOM'].cell(i,9).value)) for i in range(2,len(rows)+2)))==totals['reference_cash_cny']
for i,r in enumerate(rows,2):
    assert check['全套BOM'].cell(i,1).value==r['id']
    assert check['全套BOM'].cell(i,6).value==r['purchase_quantity']
check_import=load_workbook(HERE/'立创BOM导入-已核料号子集.xlsx',data_only=True)
assert check_import.active.max_row==6
assert {r[1].value for r in check_import.active.iter_rows(min_row=2)}=={r['lcsc_code'] for r in core}
validation=dict(date=DATE,status='passed',bom_rows=len(rows),sheets=check.sheetnames,core_import_rows=len(core),
    checks=['JSON合法与ID唯一','单行现金按MOQ和采购量重算一致','三类金额分项与总计一致','Excel重开后数量/总额一致','导入子集C编号一致，排除缺货/裸芯片/未定模块','未覆盖V1历史文件'],
    totals=totals,limitations=['未在立创云BOM上传验证列映射','没有商品结算、实物电气或装配验证','工作簿金额为静态快照，编辑后需重算'])
(HERE/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
