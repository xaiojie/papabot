"""One-sheet, whole-robot purchasing review. Does not alter the two order files."""
from collections import Counter
from decimal import Decimal
from pathlib import Path
import hashlib, json
import xlsxwriter
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OUT = HERE / '00-整机配件采购核对.xlsx'
data = json.loads((HERE/'bom-data.json').read_text())
audit = json.loads((HERE/'evidence/lcsc-match-audit.json').read_text())
cart = json.loads((HERE/'evidence/taobao-cart-audit.json').read_text())
parts = data['procurement_partition']
by = {r['id']: r for r in data['bom']}
matches = {r['id']: r for r in audit['retained_matching_rows']}
sources = ['01-立创商城-BOM导入.xlsx', '02-淘宝天猫-采购清单.xlsx']
hashes = {n: hashlib.sha256((HERE/n).read_bytes()).hexdigest() for n in sources}

def read_rows(name):
    ws = load_workbook(HERE/name).active
    return [[c.hyperlink.target if c.hyperlink else c.value for c in row] for row in list(ws)[1:]]

lcsc = {r[0]:r for r in read_rows(sources[0])}
tb = dict(zip(parts['taobao_tmall'],read_rows(sources[1]),strict=True))
assert len(lcsc)==10 and len(tb)==11
confirmed_cart = {'E01','E04','P02','P04'} | {r['id'] for r in cart['added_items']}
assert len(confirmed_cart)==9 and 'P01' not in confirmed_cart

# Concise review instructions replace old long engineering notes in the workbook.
notes = {
 'E01':'已加购1块；屏幕、旋钮和按键为同一模块，不另买。开孔按实物复测。',
 'E02':'仍需采购8个；上轮购物车未见，原链接未留存。¥80为历史截图参考；使用其中2个试单腿。',
 'E03':'已列立创BOM；价格取账户报价。附公排针未焊，不重复买；排母另列。',
 'E04':'已加购1块；VCC接3.3V。只取PWM信号，8舵机总电源走外部配电。',
 'E05':'已加购1块；接3.3V，I²S输入；拾音孔和固定位置待实测。',
 'E06':'已加购焊接款；逻辑5V供电。喇叭改焊线接螺丝端子；两路喇叭输出均不接地。',
 'E07':'已列立创BOM；自带1.25mm插头不能直插功放端子，需改焊导线，不再买PH2.0转接线。',
 'P01':'需确认≥10A、每节保护/均衡及XT30配对件。车内¥14.70普通款未核实达标，定制报价留空。',
 'P02':'已加购TYPE-C款；实际充电电流待核，目标0.4–0.5A；首版关机充电。',
 'P03':'已加购1块；只需一路舵机电源。5V/6A目标及低电量、温升待实测。',
 'P04':'原购物车已有1个；到货跳帽设5V。与舵机降压共地，两路5V输出不并联。',
 'P05':'先核对家中是否有支持数据和≥2A的A转C线；开发和充电分时共用1条。',
 'P06':'已列立创BOM；焊片式、保持开关，切断两路工作负载。面板孔/绝缘待设计。',
 'F01':'已列立创BOM；装机1个，买5个含4备件；保险值仍需结合舵机浪涌核验。',
 'F02':'已列立创BOM；这是PCB保险座，需配电载板安装，不能直接当带线保险座使用。',
 'F03':'已列立创BOM；买5个，初始装2个、备3个，数量按供电实测调整。',
 'F04':'已列立创BOM；买10个，初始装1个、备9个，最终位置与数量按电路。',
 'F06':'尚未进入下单表；需定分压电阻、RC电容、ADC脚位后生成实际料号和数量。',
 'F07':'尚未进入下单表；含本地去耦、PCA9685 OE默认禁用等小件，数值/数量随原理图确定。',
 'W01':'已列立创BOM；装机2条、买5条；配开发板随包公针，不承担8舵机公共电流。',
 'W02':'已加购1排40根；原需求30根，多10根备用，仅用于信号/小电流。',
 'W03':'已加购1米双芯，即红黑各1米。需自裁焊接；长度、压降、固定待配线。',
 'W07':'尚未进入下单表；8组舵机插座及缺少的模块排针，确认引脚顺序/接点额定；复用舵机原线。',
 'W08':'已要求电池供应商配齐公母两半；未确认包含，确认后无需另买，缺件时再补。',
 'W09':'尚未进入下单表；选小包装，按电源接头和喇叭焊点尺寸；已有可复用。',
 'W10':'尚未进入下单表；仅信号原型试接，有现成可复用；不能用细焊盘承载舵机总电流。',
 'P09':'缺配电载板设计/打样；保险座、两路降压、8路舵机配电和小件需接线落地，普通洞洞板不替代。',
 'P10':'先核对已有充电头；用于充电板5V输入，缺则买成品，具体SKU尚未选。',
 'M01':'已列立创BOM；100个为最小包装，暂用20个；长度随结构核对，不作舵盘中心螺丝。',
 'M02':'已列立创BOM；100个为最小包装，暂用12个；屏幕/外壳用孔和长度待结构核对。',
 'M03':'尚未进入下单表；普通六角螺母，捕获槽按实物设计，数量暂估。',
 'M04':'尚未进入下单表；热熔嵌件/捕获螺母只选一种，长度外径按打印材料与结构。',
 'M05':'尚未进入下单表；按最终螺丝和承压面选外径、厚度，数量暂估。',
 'M06':'尚未进入下单表；先测模块安装高度。若并入打印支架，就取消外购，不重复计。',
 'M07':'尚未进入下单表；选4个外购软脚垫或TPU打印，二选一，按脚底尺寸。',
 'M08':'尚未进入下单表；按电池包实际尺寸和托架选择长度，暂估1条。',
 'M09':'尚未进入下单表；电池缓冲和麦克风导音按用途选材，不能堵拾音孔。',
 'M10':'尚未进入下单表；小包装约10根，布局后按需补，已有可复用。',
 'M11':'缺正式打印文件和报价；主壳/前底盖、四腿、舵机座、耳饰、屏幕/电池/音频/电子件支架等整套。',
 'M12':'先做屏幕样框、单腿和公差试样，再做整套；使用已计8个舵机中的2个，不另加舵机。',
 'L01':'运费按实际订单结算；只作费用核对，非配件，不套用旧预算为实际报价。',
 'L02':'只作后续补件/返工费用核对；未发生，不提前重复购买整套配件。',
}
assert set(notes)==set(by)
status_overrides = {
 'P01':'待确认定制','E02':'待加购',
 'P05':'自备或补购','P10':'自备或补购',
 'W09':'待选小包装','W10':'待选小包装','M10':'待选小包装',
 'F06':'待电路定型','F07':'待电路定型','W07':'待电路定型',
 **{f'M{i:02d}':'待结构尺寸' for i in range(3,10)},
 'P09':'待PCB打样','M11':'待整套打印','M12':'待试样打印',
 'W08':'随电池确认','L01':'费用另计','L02':'费用另计',
}
priority = {'待确认定制':0,'待加购':1,'自备或补购':2,'待选小包装':3,'待电路定型':4,
            '待结构尺寸':5,'待PCB打样':6,'待试样打印':7,'待整套打印':8,
            '已加购':9,'已列立创BOM':10,'随电池确认':11,'费用另计':12}
installed_override = {
 'W02':'约30根','W03':'按布局裁切','P09':'1片载板','M03':'约20个','M04':'约12个',
 'M05':'约12个','M06':'约4个','M08':'约1条','M09':'按用途裁切','M10':'约10根',
 'L01':'—','L02':'—','F03':'暂2个','F04':'暂1个','M01':'暂20个','M02':'暂12个',
 'F06':'待电路','F07':'待电路','W07':'8组＋模块用针',
}
purchase_override = {
 'W02':'1排（40根）','W03':'1米双芯','P09':'1批（数量待定）','M03':'约20个/小包',
 'M04':'约12个/小包','M05':'约12个/小包','M06':'约4个或打印','M07':'4个或打印',
 'M09':'各用途小片','M10':'约10根/小包','L01':'非配件','L02':'按需补件',
 'P05':'缺则1条','P10':'缺则1个','F06':'待定','F07':'待定','W07':'待定','W08':'随包；缺再补',
}
short_names = {
 'P01':'2S动力电池包','P06':'总开关','E02':'MG90S定位舵机','F06':'电池采样/RC小件',
 'F07':'去耦/OE/接口小件','W07':'舵机接口及模块排针','P09':'载板/配电PCB',
 'M11':'外壳、腿及支架整套打印','M12':'样框/单腿/公差试样',
}
review=[]
for index,r in enumerate(data['bom']):
    rid=r['id']
    state='已加购' if rid in confirmed_cart else '已列立创BOM' if rid in parts['lcsc'] else status_overrides[rid]
    installed=installed_override.get(rid,f"{r['installed_quantity']}{r['unit']}")
    quantity=purchase_override.get(rid,f"{r['purchase_quantity']}{r['unit']}")
    spec=r['model']
    url=''; channel='待选渠道'; cost=None
    if rid in tb:
        source=tb[rid]
        assert source[5]==r['url'] or (not source[5] and not r['url'])
        spec=source[1]
        channel='淘宝/天猫 · '+source[4]
        url=source[5] or ''
        quantity=purchase_override.get(rid,source[2])
        if source[3] is not None:
            cost=float(Decimal(str(source[3]))*Decimal(str(r['purchase_quantity'])))
    elif rid in matches:
        match=matches[rid]; source=lcsc[match['matched_code']]
        assert source[1]==match['matched_model'] and source[2]==int(match['purchase_quantity'])
        spec=f"{match['matched_model']}\n{match['matched_code']}"
        channel='立创商城 · '+match['matched_code']; url=r['url']
        quantity=f"{source[2]}{r['unit']}"; cost=match['line_cash_cny']
    elif rid in {'P09'}: channel='嘉立创PCB · 待设计报价'
    elif rid in {'M11','M12'}: channel='3D打印 · 待报价'
    elif rid=='W08': channel='电池供应商 · 随包确认'
    elif rid in {'L01','L02'}: channel='非配件 · 费用核对'
    elif rid in {'P05','P10'}: channel='先核对已有配件'
    review.append(dict(id=rid,source_order=index,status=state,name=short_names.get(rid,r['item']),
                       spec=spec,installed=installed,purchase=quantity,channel=channel,url=url,cost=cost,note=notes[rid]))
review.sort(key=lambda r:(priority[r['status']],r['source_order']))
known_total=sum(Decimal(str(r['cost'])) for r in review if r['cost'] is not None)
assert known_total==Decimal('369.64')
headers=['当前状态','配件','型号 / 规格','整机用量','本次采购量','渠道 / 商品链接','已知金额(元)','还需核对什么']

wb=xlsxwriter.Workbook(OUT,{'strings_to_urls':False})
ws=wb.add_worksheet('整机采购核对')
ws.hide_gridlines(2); ws.set_zoom(85); ws.freeze_panes(1,2)
base={'font_name':'PingFang SC','font_size':11,'valign':'vcenter','text_wrap':True}
head=wb.add_format({**base,'bold':True,'font_color':'white','bg_color':'#365A50'})
formats={}
for band,bg in enumerate(['#FFFFFF','#F7F8FA']):
    for kind,more in {'text':{},'money':{'num_format':'0.00'},'link':{'font_color':'#176B87','underline':1}}.items():
        formats[band,kind]=wb.add_format({**base,'bg_color':bg,**more})
colors={'已加购':('#E5F3EA','#1D6F42'),'已列立创BOM':('#E7EFFB','#2D5390'),
        '随电池确认':('#F0F1F3','#555D68'),'费用另计':('#F0F1F3','#555D68')}
state_formats={s:wb.add_format({**base,'bold':True,'bg_color':colors.get(s,('#FFF0D9','#8C5300'))[0],
                               'font_color':colors.get(s,('#FFF0D9','#8C5300'))[1]}) for s in priority}
ws.write_row(0,0,headers,head); ws.set_row(0,32)
for col,width in enumerate([18,25,37,14,18,27,16,49]): ws.set_column(col,col,width)
ws.write_comment(0,0,'状态依据2026-09-08已保存的采购表与购物车核验，不代表下单或收货。此总表用于核对，不能整体导入立创BOM。')
ws.write_comment(0,6,'按本次采购量的小计。立创来自账户匹配报价，淘宝来自已选SKU价格；舵机为历史截图参考。未报价项目留空，旧预算占位不计入。')
for row,r in enumerate(review,1):
    band=(row-1)%2
    vals=[r[k] for k in ['status','name','spec','installed','purchase','channel','cost','note']]
    ws.set_row(row,62 if len(r['note'])<64 else 76)
    for col,val in enumerate(vals):
        fmt=state_formats[r['status']] if col==0 else formats[band,'money' if col==6 else 'text']
        if col==5 and r['url']: ws.write_url(row,col,r['url'],formats[band,'link'],r['channel'])
        elif val is None: ws.write_blank(row,col,None,fmt)
        else: ws.write(row,col,val,fmt)
    if r['id']=='M11':
        ws.write_comment(row,2,'整套打印的计划构成（不另计费，数量未冻结）：\n'+'\n'.join(f'{p[1]} × {p[2]}' for p in data['planned_print_parts']),{'width':420,'height':520})
ws.autofilter(0,0,len(review),7)
sum_row=len(review)+1
sumfmt=wb.add_format({**base,'bold':True,'bg_color':'#E7EDE9','num_format':'0.00'})
ws.set_row(sum_row,46)
ws.write(sum_row,5,'已知商品金额（筛选小计）',sumfmt)
ws.write_formula(sum_row,6,f'=SUBTOTAL(109,G2:G{len(review)+1})',sumfmt,float(known_total))
ws.write(sum_row,7,'金额空白项、运费及制作费用未含；已加购≠已购买。',sumfmt)
ws.set_row(sum_row+2,42)
ws.write(sum_row+2,1,'使用说明',formats[0,'text'])
ws.write(sum_row+2,2,'优先看表头的待办项目；筛选“当前状态”核对缺口。',formats[0,'text'])
ws.write(sum_row+2,5,'核验记录：2026-09-08',formats[0,'text'])
ws.write(sum_row+2,7,'本表只供整机核对；立创下单仍用原01文件。',formats[0,'text'])
ws.set_row(sum_row+3,46)
ws.write(sum_row+3,1,'装配工具另核对',formats[0,'text'])
ws.write(sum_row+3,2,'电烙铁、焊锡/助焊剂、万用表、剥线/剪线工具、螺丝刀、热缩工具。',formats[0,'text'])
ws.write(sum_row+3,7,'已有可复用；工具不计入整机配件数量或已知金额。',formats[0,'text'])
ws.set_landscape(); ws.set_paper(9); ws.fit_to_pages(1,0)
ws.set_margins(.25,.25,.35,.35); ws.repeat_rows(0); ws.set_footer('&C整机采购核对 · 2026-09-08  |  第 &P / &N 页')
wb.close()

# Verify source consistency, quantities, hyperlinks, missing-price handling and cached total.
out=load_workbook(OUT); check=out.active
assert len(out.worksheets)==1 and check.max_column==8 and not check.merged_cells.ranges
assert len(review)==42 and {r['id'] for r in review}==set(by)
assert [c.value for c in check[1]]==headers
for ri,r in enumerate(review,2):
    expected=[r[k] for k in ['status','name','spec','installed','purchase','channel','cost','note']]
    for ci,v in enumerate(expected,1): assert check.cell(ri,ci).value==v,(r['id'],ci)
    if r['url']: assert check.cell(ri,6).hyperlink.target==r['url']
    else: assert check.cell(ri,6).hyperlink is None
    if r['id'] not in matches and (r['id'] not in tb or r['id']=='P01'): assert check.cell(ri,7).value is None
assert load_workbook(OUT,data_only=True).active.cell(sum_row+1,7).value==369.64
assert all(hashlib.sha256((HERE/n).read_bytes()).hexdigest()==h for n,h in hashes.items())
counts=dict(Counter(r['status'] for r in review))
result={'date':'2026-09-08','file':OUT.name,'sheets':1,'columns':8,'review_rows':42,
        'status_counts':counts,'known_priced_rows':20,'known_subtotal_cny':float(known_total),
        'all_42_engineering_records_covered_once':True,'purchase_inputs_unchanged':True,
        'purchase_input_sha256':hashes,'cart_snapshot_reused_not_refreshed':True,
        'empty_amount_means_unquoted_not_free':True,'new_orders_or_cart_changes':False,
        'status':'passed'}
(HERE/'review-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
(HERE/'review-data.json').write_text(json.dumps({'meta':result,'rows':review},ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
