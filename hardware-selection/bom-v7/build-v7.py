"""Simple procurement sheets from verified cart rows, 2026-09-08."""
from pathlib import Path
from decimal import Decimal
from collections import Counter
import json, shutil, hashlib, copy
import xlsxwriter
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
OLD=HERE.parent/'bom-v6'
ROOT=HERE.parent.parent
def dump(path,data): path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
old=json.loads((OLD/'bom-data.json').read_text())
old_review=json.loads((OLD/'review-data.json').read_text())['rows']
by={r['id']:r for r in old_review}

# Each line is one SKU observed in the refreshed cart, quantity 1 package.
# Fields: procurement id, engineering requirements, item id, SKU id, short name,
# exact cart variant, package contents, displayed package price, shop, use/limitation.
new=[
 ['W09',['W09'],'615600154622','4334209992605','绝缘热缩管','127只 黑色 盒装','1盒（127段）',8.62,'touglesy旗舰店','按焊点选择管径后收缩；余量作备用。'],
 ['W10',['W10'],'616867442654','5864266635679','信号试接洞洞板','4x6 双面喷锡洞洞板（5块）','1包（5块）',6.35,'达科峰旗舰店','FR-4双面喷锡，40×60mm；暂用1块，其余备用。只用于信号试接，舵机总电流走独立配电。'],
 ['W07',['W07'],'606726184415','4430371271852','舵机接口及模块排针','单排黑--镀金（10只）','1包（10条40P）',4.52,'zave旗舰店','2.54mm单排直针，可裁成3P×8及模块短针；复用模块自带排针。不能用单个接点承载8舵机总电流。'],
 ['M06',['M06'],'704528035119','5130695735303','尼龙绝缘垫柱试装盒','180PCS白色','1盒（180件混合）',11.16,'深圳市壹优创电子','M3尼龙螺丝/螺母/垫柱混合盒；高度按实物试配。180件是套装总件数，不是180根垫柱。'],
 ['M03',['M03'],'673051124656','5054000291214','M2六角螺母','M2【20个】316不锈钢','1包（20个）',0.96,'亿豪紧固件企业店','普通公制右牙，配M2机牙螺丝；捕获槽按实物留公差。'],
 ['M04',['M04'],'673051124656','5054000291216','M3六角螺母','M3【50个】316不锈钢','1包（50个）',1.19,'亿豪紧固件企业店','选择普通螺母捕获槽方案；暂用约12个，不再另买热熔铜嵌件。'],
 ['M05',['M05'],'575265264956','5076932416603','M3平垫片','M3*6*0.8[200个]','1包（200个）',2.77,'固万基官方旗舰店','M3、外径6mm、厚0.8mm；暂用约12个。此SKU为200片小袋，余件备用，非200台用量。'],
 ['M07',['M07'],'771706659250','5290081145877','自粘硅胶软脚垫','灰色带胶 直径12mm*2mm/50只装','1包（50只）',6.83,'硅胶制品厂','暂用4只，余46只备用；脚底粘贴面按Ø12mm设计，实物测试抓地力。'],
 ['M08',['M08'],'773700829610','5466028503337','电池反扣魔术贴束带','黑色1cm宽*15cm长【2条】','1包（2条）',1.60,'粘扣带制造','10×150mm，暂用1条另1条备用；按成品电池和托架试配，不勒压软包。'],
 ['M10',['M10'],'676815692302','5034408146913','小号尼龙扎带','国标3*100黑【宽2.5mm】100条','1包（100条）',3.50,'浙江力泰塑料工厂店','实际宽2.5mm、长100mm；暂用约10条，只束线，不直接捆压电芯。'],
 ['M09',['M09'],'645815964629','4660038213867','自粘EVA缓冲/密封泡棉','20mm宽 10米长 1mm厚','1卷',3.00,'敏佰旗舰店','裁短条作电池缓冲；麦克风只裁环形密封垫，中心留通孔。不是透声网或覆盖拾音孔的海绵。'],
 ['P10',['P10'],'39662398527','6289267198584','绿联USB-A充电头','17W丨白色双口USB丨小巧低温|直插','1个',27.90,'绿联数码旗舰店','CD104，固定5V，单口最大2.4A、双口总3.4A。专用一口接充电板；裸头不含线。'],
 ['F06A',['F06'],'982371708131','5939246944184','电池采样上臂电阻','1/4W 金属膜电阻1% 100K （100个）','1包（100个）',2.54,'telesky旗舰店','原型上臂100kΩ；与33kΩ下臂组成分压，8.4V约降至2.08V。装机暂1个；测阻并校准ADC。'],
 ['F06B',['F06'],'982371708131','5939246944172','电池采样下臂电阻','1/4W 金属膜电阻1% 33K （100个）','1包（100个）',2.54,'telesky旗舰店','原型下臂33kΩ接地；装机暂1个，ADC引脚和衰减档随原理图/固件确定。'],
 ['F07A',['F07'],'982371708131','5939246944162','PCA9685 OE上拉电阻','1/4W 金属膜电阻1% 10K （100个）','1包（100个）',2.54,'telesky旗舰店','OE拟通过10kΩ上拉3.3V，默认禁止PWM；接线前核实模块原有上下拉，必要时改板；方案未实施。'],
 ['F06C',['F06','F07'],'609963289065','4286902992281','采样滤波及本地去耦电容','104P 0.1UF 50V 脚距5.08MM ±20%（50个）','1包（50个）',3.64,'jccon旗舰店','100nF直插独石电容，1个用于采样RC，其余用于需要补充的本地去耦。两项需求共用一包，不重复购买。'],
]
old_ids=['E01','E02','E04','E05','E06','P01','P02','P03','P04','W02','W03']
shop_fix={'E02':'玖訫科技','P01':'明盛电子商城全新原装'}
actual=[]
for rid in old_ids:
    r=copy.deepcopy(by[rid])
    if rid=='E02':
        r.update(url='https://item.taobao.com/item.htm?id=1005683789132&skuId=5999915261650',cost=88.0,
                 spec='MG90S舵机（全金属齿轮）180度',note='已核购物车数量8，单价11元；每腿2个；试单腿复用其中2个，不另买。')
    if rid=='P01':
        r.update(url='https://item.taobao.com/item.htm?id=966262592635&skuId=5908158073708',
                 spec='车内普通款：电压7.4V 长53mm宽34mm厚10mm；定制规格另见03文件',
                 note='原店已由用户询问并回复“没有”，排除该供应商；车内14.70元普通款勿下单。新候选见04换店文件，整包持续≥10A及保护/均衡均待答复，定制价格留空。')
    qty=8 if rid=='E02' else 1
    price=None if rid=='P01' else float(Decimal(str(r['cost']))/qty)
    actual.append(dict(id=rid,requirements=[rid],name=r['name'],sku=r['spec'],quantity=qty,
                       pack='8个' if rid=='E02' else r['purchase'],price=price,
                       shop=shop_fix.get(rid,r['channel'].split(' · ')[-1]),url=r['url'],
                       note=r['note'],status='原店无货·待换店' if rid=='P01' else '已加购',
                       cart_observed_unit_price=14.7 if rid=='P01' else price,
                       added_this_turn=False))
for rid,req,item,sku,name,spec,pack,price,shop,note in new:
    actual.append(dict(id=rid,requirements=req,name=name,sku=spec,quantity=1,pack=pack,price=price,
                       shop=shop,url=f'https://item.taobao.com/item.htm?id={item}&skuId={sku}',note=note,
                       status='已加购',cart_observed_unit_price=price,added_this_turn=True))
assert len(actual)==27
known=sum(Decimal(str(r['price']))*r['quantity'] for r in actual if r['price'] is not None)
added=sum(Decimal(str(r['price'])) for r in actual if r['added_this_turn'])
cart_display_rows=known+Decimal('14.70')
shutil.copy2(OLD/'01-立创商城-BOM导入.xlsx',HERE/'01-立创商城-BOM导入.xlsx')
shutil.copy2(OLD/'evidence/lcsc-match-audit.json',HERE/'evidence/lcsc-match-audit.json')
for name in ['README.md','START-HERE.txt','design-selection.json']:
    dst=HERE/'history'/name
    if not dst.exists(): shutil.copy2(ROOT/name,dst)
for name in ['power-selection-2s.md','power-selection-2s.json']:
    dst=HERE/'history'/name
    if not dst.exists(): shutil.copy2(HERE.parent/name,dst)

audit={'date':'2026-09-08','source':'Chrome购物车可见DOM；逐行标题、SKU链接、数量输入框和价格',
       'authorization':'补齐Excel缺项并加入淘宝购物车；Type-C线取消，PCB和3D打样暂缓；8舵机已由用户加购',
       'cart_before_rows':11,'cart_after_rows':27,'added_sku_rows':16,
       'all_quantities_verified':True,'servo_quantity':8,'other_row_quantities':1,
       'existing_sku_quantities_preserved':True,'new_cart_line_display_subtotal_cny':float(added),
       'all_cart_line_display_subtotal_including_unqualified_battery_cny':float(cart_display_rows),
       'selected_procurement_reference_excluding_custom_battery_cny':float(known),
       'price_note':'逐行展示优惠价，仅作参考。单品券/首单礼金可能互斥；未进入结算，不代表整车实付或含运费。',
       'ordinary_battery_qualified':False,'ordinary_battery_checkbox_observed':False,
       'order_placed':False,'seller_messages_sent':False,'online_lcsc_bom_changed':False,
       'items':actual}
dump(HERE/'evidence/taobao-cart-audit.json',audit)

# Expand grouped circuit requirements into separately buyable SKUs; one shared capacitor bag.
battery_provisional={
    'status':'暂定·待尺寸报价',
    'shop':'锂电池总汇店',
    'url':'https://item.taobao.com/item.htm?id=986003989892',
    'spec':'2S1P 7.4V/1000mAh；整包持续≥10A；保护/均衡；XT30配套',
    'note':'用户反馈新店可以做，暂定此店。目标含保护板包装≤75×40×20mm、≤80g；18AWG约10cm及XT30配对线束、共口充放电/支持1A充电，待完整规格与报价确认。未加购定制款；勿买车内原店普通电池。',
    'price':None,
    'confirmation_source':'用户转述第一家店家说可以；尺寸、重量、线材、报价和交期未获最终回复',
}
review=[]
for r in actual:
    oldr=by.get(r['id'])
    installed=oldr['installed'] if oldr else ('暂1个' if r['id']!='F06C' else '采样1个＋按需去耦')
    if r['id']=='M06': installed='约4个垫柱'
    if r['id']=='W07': installed='3P×8＋短排针'
    review.append(dict(id=r['id'],requirements=r['requirements'],status=r['status'],name=r['name'],
                       spec=r['sku'],installed=installed,purchase=r['pack'],channel='淘宝 · '+r['shop'],
                       url=r['url'],cost=None if r['price'] is None else float(Decimal(str(r['price']))*r['quantity']),note=r['note']))
for r in review:
    if r['id']=='P01':
        r.update(status=battery_provisional['status'],spec=battery_provisional['spec'],
                 channel='淘宝 · '+battery_provisional['shop'],url=battery_provisional['url'],
                 cost=None,note=battery_provisional['note'])
for r in old_review:
    if r['status']=='已列立创BOM':
        review.append({**r,'requirements':[r['id']]})
    elif r['id'] in ['P05','P09','M11','M12','W08','L01','L02']:
        r=copy.deepcopy(r); r['requirements']=[r['id']]
        if r['id']=='P05': r.update(status='本次不采购',purchase='0条',note='用户明确不需要采购Type-C数据/充电线。',channel='用户取消本次采购')
        if r['id'] in ['P09','M11','M12']: r.update(status='暂缓',purchase='本次0',note='按用户要求暂缓，设计尚未完成；本轮不选商家、不加购、不计预算。')
        if r['id']=='W08': r.update(status='随电池定制',purchase='随电池配1对',note='并入定制电池订单，要求电池端已焊XT30＋机器端配对半边，18AWG引线；本次不重复加购。')
        review.append(r)
priority={'暂定·待尺寸报价':0,'随电池定制':1,'已加购':2,'已列立创BOM':3,'本次不采购':4,'暂缓':5,'费用另计':6}
review.sort(key=lambda r:(priority[r['status']],r['id']))
assert len(review)==44
assert set(x for r in review for x in r['requirements'])==set(by)
total=sum(Decimal(str(r['cost'])) for r in review if r['cost'] is not None)

def workbook(path,sheet,headers,rows,widths,link_col,money_cols,states=False,total_col=None):
    wb=xlsxwriter.Workbook(path,{'strings_to_urls':False})
    ws=wb.add_worksheet(sheet); ws.hide_gridlines(2); ws.freeze_panes(1,2); ws.set_zoom(85)
    base={'font_name':'PingFang SC','font_size':11,'valign':'vcenter','text_wrap':True}
    head=wb.add_format({**base,'bold':True,'font_color':'white','bg_color':'#365A50'})
    fmts={}
    for band,bg in enumerate(['#FFFFFF','#F5F7F7']):
        for kind,more in {'text':{},'money':{'num_format':'0.00'},'link':{'font_color':'#166884','underline':1}}.items():
            fmts[band,kind]=wb.add_format({**base,'bg_color':bg,**more})
    warn=wb.add_format({**base,'bg_color':'#FFF0D9','font_color':'#875000','bold':True})
    ws.write_row(0,0,headers,head); ws.set_row(0,32)
    for i,w in enumerate(widths): ws.set_column(i,i,w)
    for ri,(vals,url) in enumerate(rows,1):
        ws.set_row(ri,72 if len(str(vals[-1]))>65 else 60)
        for ci,val in enumerate(vals):
            fmt=fmts[(ri-1)%2,'money' if ci in money_cols else 'text']
            if states and ci==0 and val not in ['已加购','已列立创BOM']: fmt=warn
            if ci==link_col and url: ws.write_url(ri,ci,url,fmts[(ri-1)%2,'link'],str(val))
            elif val is None: ws.write_blank(ri,ci,None,fmt)
            else: ws.write(ri,ci,val,fmt)
    ws.autofilter(0,0,len(rows),len(headers)-1)
    if total_col is not None:
        sr=len(rows)+1; sf=wb.add_format({**base,'bold':True,'bg_color':'#E6EEE9','num_format':'0.00'})
        ws.set_row(sr,50); ws.write(sr,total_col-1,'已知商品参考小计',sf)
        letter=xlsxwriter.utility.xl_col_to_name(total_col)
        ws.write_formula(sr,total_col,f'=SUBTOTAL(109,{letter}2:{letter}{len(rows)+1})',sf,float(total))
        ws.write(sr,len(headers)-1,'不含定制电池、运费、打样；优惠以实际结算为准。已加购不等于已下单。',sf)
    ws.write_comment(0,0,'2026-09-08核验。已加购表示购物车已存在，不代表已下单或验证整机可运行。')
    for c in money_cols: ws.write_comment(0,c,'按已选SKU展示价格记录；首单礼金、单品券等可能不能叠加，最终以实际结算为准。空白为未报价，不是免费。')
    ws.set_landscape();ws.set_paper(9);ws.fit_to_pages(1,0);ws.repeat_rows(0);ws.set_margins(.25,.25,.35,.35)
    ws.set_footer('&C桌面助手采购 V7 · 2026-09-08  |  &P / &N')
    wb.close()

tbrows=[]
for r in sorted(actual,key=lambda r:(r['id']=='P01',r['id'])):
    if r['id']=='P01':
        tbrows.append(([r['name'],battery_provisional['spec'],r['pack'],None,battery_provisional['shop'],
                        '打开定制入口',battery_provisional['note']],battery_provisional['url']))
        continue
    q=f"{r['quantity']}个" if r['id']=='E02' else r['pack']
    note=('暂勿下单普通款；' if r['id']=='P01' else '已加购；')+r['note']
    tbrows.append(([r['name'],r['sku'],q,r['price'],r['shop'],'打开已选商品',note],r['url']))
workbook(HERE/'02-淘宝天猫-采购清单.xlsx','淘宝采购',['配件','购买规格','采购量','参考单价(元)','店铺','商品链接','备注'],tbrows,[24,49,22,16,26,19,63],5,[3])
rrows=[([r[k] for k in ['status','name','spec','installed','purchase','channel','cost','note']],r['url']) for r in review]
workbook(HERE/'00-整机配件采购核对.xlsx','整机采购核对',['当前状态','配件','型号 / 规格','整机用量','本次采购量','渠道 / 商品链接','参考金额(元)','核对说明'],rrows,[18,25,43,19,23,27,17,64],5,[6],True,6)

data={'schema_version':7,'date':'2026-09-08','currency':'CNY','scope':'当前配件全部暂定；取消Type-C线，暂缓PCB/3D；新电池商家表示可以做，待尺寸和报价确认',
      'battery_supplier_provisional':battery_provisional,
      'source_engineering_bom':'../bom-v6/bom-data.json','engineering_requirement_count':42,
      'bom':review,'taobao_cart_items':actual,
      'procurement_totals':{'lcsc_reference_cny':111.37,'taobao_reference_excluding_custom_battery_cny':float(known),
                            'new_16_sku_reference_cny':float(added),'known_reference_total_cny':float(total)},
      'procurement_partition':{'lcsc':[r['id'] for r in review if r['status']=='已列立创BOM'],
                               'taobao_tmall':[r['id'] for r in actual],'no_extra_official_purchase':True,
                               'cancelled':['P05'],'deferred':['P09','M11','M12'],'battery_bundle':['P01','W08'],'fees':['L01','L02']},
      'production_design_complete':False,'order_placed':False}
dump(HERE/'bom-data.json',data);dump(HERE/'review-data.json',{'rows':review})

# Validate against source identities, independent price sums, and the shipped files.
keys=[(r['url'].split('id=')[1].split('&')[0],r['url'].split('skuId=')[-1] if 'skuId=' in r['url'] else '') for r in actual]
assert len(set(keys))==27
assert len([r for r in actual if r['added_this_turn']])==16
assert next(r for r in actual if r['id']=='E02')['quantity']==8
assert all(r['quantity']==1 for r in actual if r['id']!='E02')
assert known==Decimal('355.93') and added==Decimal('89.66') and total==Decimal('467.30')
assert next(r for r in actual if r['id']=='P01')['price'] is None
assert len([r for r in actual if r['id']=='F06C'])==1
assert not any(r['id'] in ['P05','P09','M11','M12','W08','L01','L02'] for r in actual)
hash01=hashlib.sha256((HERE/'01-立创商城-BOM导入.xlsx').read_bytes()).hexdigest()
assert hash01=='83caaaf752e2bd71f4888ba88c8e093c8f1aba2576d3a6be60e1d3cb982fb767'
for filename,count,cols in [('01-立创商城-BOM导入.xlsx',10,3),('02-淘宝天猫-采购清单.xlsx',27,7),('00-整机配件采购核对.xlsx',44,8)]:
    w=load_workbook(HERE/filename); s=w.active
    assert len(w.worksheets)==1 and s.max_column==cols and not s.merged_cells.ranges
    assert s.auto_filter.ref==f'A1:{xlsxwriter.utility.xl_col_to_name(cols-1)}{count+1}'
    assert s.freeze_panes
for fn,rows,lc in [('02-淘宝天猫-采购清单.xlsx',tbrows,5),('00-整机配件采购核对.xlsx',rrows,5)]:
    s=load_workbook(HERE/fn).active
    for i,(vals,url) in enumerate(rows,2):
        for j,val in enumerate(vals,1): assert s.cell(i,j).value==val,(fn,i,j)
        if url: assert s.cell(i,lc+1).hyperlink.target==url
assert load_workbook(HERE/'00-整机配件采购核对.xlsx',data_only=True).active.cell(46,7).value==float(total)
validation={'status':'passed','date':'2026-09-08','one_sheet_each':True,'lcsc_rows':10,'taobao_rows':27,'review_rows':44,
            'all_42_requirements_covered':True,'all_27_cart_sku_quantities_verified':True,'new_sku_rows':16,
            'servo_quantity':8,'duplicate_cart_sku_rows':0,'shared_capacitor_counted_once':True,
            'excluded_cable_and_services_absent_from_taobao_purchase':True,'lcsc_unchanged_sha256':hash01,
            'known_reference_total_cny':float(total),'new_reference_subtotal_cny':float(added),
            'battery_custom_quote_included':False,'order_placed':False}
dump(HERE/'validation.json',validation)
print(json.dumps(validation,ensure_ascii=False,indent=2))
