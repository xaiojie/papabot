"""Export the complete current sourcing list in the user's LCSC column order.

This is a dated sourcing draft, not a completed PCB manufacturing BOM.
Rebuilding only writes inside bom-v3. It never touches the online BOM.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import Counter
import json
import csv
import copy
import xlsxwriter
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DATE = '2026-09-08'
data = copy.deepcopy(json.loads((ROOT/'hardware-selection/bom-v2/bom-data.json').read_text()))
rows = data['bom']
by_id = {r['id']:r for r in rows}

def rounded(value):
    return float(Decimal(str(value)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))

by_id['E04'].update(model='PCA9685 16路PWM成品模块',
    unit_price_or_allowance_cny=8.9,price_basis='本轮网页价',line_cash_cny=8.9,
    channel='淘宝 · 汉丽电子科技',url='https://item.taobao.com/item.htm?id=967408745108',
    selection_status='新补商品；推荐样机候选，待电气/尺寸核对',observed_stock='有货',observed_moq=1,
    notes='HANLI品牌，16路PWM/IIC成品板；详情页面8.90元、快递5元（运费计入L01统一预留，不重复加）。不通过此板承担8舵机总电流。OE、逻辑3.3V及V+独立接线须按实物核对。',
    mechanical_constraints='商家页面未提供可读尺寸表，板框/高件/孔位待核。')
by_id['W02'].update(item='母对母信号杜邦线，每包10根',model='FIT0030',
    unit='包',installed_quantity=3,purchase_quantity=3,unit_price_or_allowance_cny=6,
    line_cash_cny=18,price_basis='本轮网页价',channel='DFRobot 官网',lcsc_code='',
    url='https://www.dfrobot.com.cn/goods-366.html',observed_stock='有货',observed_moq=1,
    selection_status='新补商品；母母头和长度已核',
    notes='每包10根、20cm、母对母；3包共30根用于信号和小电流。官网标注电流<1000mA，不作舵机/功放/电池主电源线；机内长度后续按线束表整理。替换V2未核头型的LCKFB线。',
    mechanical_constraints='20cm成品线，两端母头；PCB端需2.54mm公排针。')
by_id['W03'].update(item='红黑18AWG硅胶电源线套装，各1米',model='FIT0585',
    installed_quantity=1,purchase_quantity=1,unit='套',unit_price_or_allowance_cny=10,
    line_cash_cny=10,price_basis='本轮网页价',channel='DFRobot 官网',
    url='https://www.dfrobot.com.cn/goods-1770.html',observed_stock='有货',observed_moq=1,
    selection_status='新补小包装商品；用于主干和短分支',
    notes='红1m＋黑1m，0.75mm²多股镀锡铜/硅胶；替换V2红黑18AWG及红黑20AWG四行预留。8个舵机自带线复用；本套用于电池、两路降压和配电板的短连线。长度/压降/成束温升待布局验证，不将网页22A宣称视为本机额定。',
    mechanical_constraints='接头焊杯、模块焊盘、弯曲和固定空间需容纳18AWG。')
rows = [r for r in rows if r['id'] not in {'W04','W05','W06'}]

brands = {'E01':'TOUGLESY','E03':'立创开发板','E04':'HANLI','E05':'其它品牌',
    'E06':'DFRobot','E07':'XHXDZ','P02':'TELESKY','P03':'DFRobot','P04':'DFRobot',
    'F01':'Littelfuse','F02':'Littelfuse','F03':'KNSCHA','F04':'JIERR','W01':'JILN',
    'W02':'其它品牌','W03':'其它品牌','M01':'顺天五金','M02':'顺天五金'}
exact_models = {'E01':'ST7789 TFT＋EC11（2.4英寸）','E02':'MG90S 180度（用户既定店铺款）',
    'E04':'PCA9685 16路PWM成品模块','E06':'DFR0954','P02':'MK021163 TYPE-C 2S8.4V',
    'F01':'029707.5WXNV','F02':'01530008Z','F03':'SHW1000UF10V8X12','F04':'ERB16V470M8X12',
    'M01':'PM2X6','M02':'PM3X8'}
non_import={'W08','P09','M11','M12','L01','L02'}
pending_model={'P01','P05','P06','F05','F06','F07','W07','W09','W10','P10',
    'M03','M04','M05','M06','M07','M08','M09','M10'}
for r in rows:
    r['original_brand']=brands.get(r['id'],'')
    r['original_package']=''
    r['original_model']=exact_models.get(r['id'],r['model']) if r['id'] not in pending_model|non_import else ''
    r['spare_quantity']=max(0,r['purchase_quantity']-r['installed_quantity'])
    r['export_to_full_material_sheet']=r['id'] not in non_import
    r['exact_purchase_sku_available']=bool(r['original_model'])
    r['line_cash_cny']=rounded(Decimal(str(r['unit_price_or_allowance_cny']))*Decimal(str(r['purchase_quantity'])))
    r['price_observation_date']=DATE if r['price_basis']=='本轮网页价' else None
    r['notes']+='；本版按1套样机计量。'

materials=[r for r in rows if r['export_to_full_material_sheet']]
lcsc=[r for r in materials if r['lcsc_code']]
pending=[r for r in materials if not r['original_model']]
excluded=[r for r in rows if not r['export_to_full_material_sheet']]
total=rounded(sum(Decimal(str(r['line_cash_cny'])) for r in rows))
lcsc_total=rounded(sum(Decimal(str(r['line_cash_cny'])) for r in lcsc))
groups={basis:rounded(sum(Decimal(str(r['line_cash_cny'])) for r in rows if r['price_basis']==basis))
    for basis in ['本轮网页价','历史用户截图价','预算占位','已含在其他行']}

headers=['原型号','原品牌','原封装','原参数','原C编码','单套用量','备损量',
    '单位','物料编号','配件名称','选型状态','采购渠道','商品链接',
    '已查MOQ','建议采购数量','参考单价/预留(元)','参考现金(元)','价格依据','备注']

def values(r):
    params=f"{r['item']}；{r['model']}；单位：{r['unit']}"
    if not r['original_model']:
        params='待定规格，仅列需求：'+params
    return [r['original_model'],r['original_brand'],r['original_package'],params,
        r['lcsc_code'],r['installed_quantity'],r['spare_quantity'],r['unit'],r['id'],
        r['item'],r['selection_status'],r['channel'],r['url'],r['observed_moq'],
        r['purchase_quantity'],r['unit_price_or_allowance_cny'],r['line_cash_cny'],
        r['price_basis'],r['notes']+' '+r['mechanical_constraints']]

guide=[
    ['日期',DATE],
    ['导入对象',f'首个工作表“BOM导入_全量”含{len(materials)}项实物需求；完整总表{len(rows)}项。'],
    ['当前完成度',f'已有具体型号或既定店铺款{len(materials)-len(pending)}项；另有{len(pending)}项需定制/选商品/电路或结构设计。不是全部选型已完成的最终采购单。'],
    ['网页操作','用户自行导入新版BOM。本轮未修改当前云BOM、未上传文件、未加入购物车。'],
    ['步骤1','在立创BOM导入Excel，选第一个“BOM导入_全量”工作表；第1行作为表头。'],
    ['步骤2','映射：原型号→原型号，原品牌→原品牌，原封装→原封装，原参数→原参数，原C编码→原C编码，单套用量→单套用量，备损量→备损量。'],
    ['步骤3','采购套数设1。已查MOQ的多买余量已计入备损量，不要再把建议采购数量叠加到单套用量。'],
    ['数量示例','M2×6螺丝：单套用量20、备损量80，合计100；1000µF电容：单套2、备损3，合计5。'],
    ['备损量口径','本文件备损量含MOQ产生的多买余量及备件，均为绝对数量。导入界面若采用百分比，勿映射本列；改用建议采购数量作为单套用量并置备损0，仅用于1套采购。'],
    ['单位','有的行按个，有的按包/套。FIT0030的单套用量3表示3包共30根，不是3根。页面原参数和备注均保留包装信息。'],
    ['无法完全匹配属于预期','淘宝/DFRobot外采没有已核立创C编号；待选项原型号及C编号留空，需求放在原参数。不能将自动推荐的裸芯片直接替代成品模块。'],
    ['封装','没有经过核实的封装字段留空；外形/尺寸放在原参数或备注，不把商品包装“袋装”当电气封装。'],
    ['只导入有C编号物料',f'可用单独的“立创BOM导入-全部已知C编号V3.xlsx”，含{len(lcsc)}项及对应条件，参考{lcsc_total:.2f}元。包括MRO；若云BOM不支持MRO自动匹配，则按商品链接在工业品核对。'],
    ['采购候选的限制','保险/保险座取决于载板，螺丝长度取决于结构。查到型号或C编号不等于整机适配、库存结算、运动/温升已经验证。'],
    ['费用与包内件','PCB/打印服务、运费/返工预留以及电池已含XT30配对件放在单独工作表，避免作为电子料自动匹配或重复采购。'],
    ['完整参考现金',total],
    ['网页观察价',groups['本轮网页价']],
    ['历史截图价',groups['历史用户截图价']],
    ['仍为预算占位',groups['预算占位']],
    ['总预算条件','两路仍按各1块DFR0753备选计价；新的总预算上限未确认。未定价格不表示已获供应商报价。'],
    ['后续收尾','先确定电源开关/载板原理图、定制电池供货与充电限流，再定端子/线束和结构紧固。更新同一物料编号行，补真实型号/C编号后重新导出。'],
    ['换套数','本文件按1套样机+本次备件/MOQ生成；增加套数需按实际需求和MOQ重新计算，不能直接沿用所有备损量。'],
    ['数据计算','导入工作表全部是静态数值，避免上传器把公式读成空值。修改数量或单价后需重算金额；生成器可复现本版。'],
]

def make_book(path, primary, full):
    wb=xlsxwriter.Workbook(path,{'strings_to_urls':False})
    fmt_head=wb.add_format({'bold':True,'font_color':'#FFFFFF','bg_color':'#365A50','text_wrap':True,'valign':'vcenter','font_name':'PingFang SC'})
    fmt_text=wb.add_format({'font_name':'PingFang SC','font_size':10,'text_wrap':True,'valign':'top'})
    fmt_num=wb.add_format({'font_name':'PingFang SC','num_format':'0.000000','valign':'top'})
    fmt_money=wb.add_format({'font_name':'PingFang SC','num_format':'0.00','valign':'top'})
    fmt_pending=wb.add_format({'bg_color':'#FFF2CC','font_color':'#7F6000','text_wrap':True,'valign':'top','font_name':'PingFang SC'})
    fmt_link=wb.add_format({'font_color':'#176B87','underline':1,'text_wrap':True,'valign':'top','font_name':'PingFang SC'})
    def table(name, rs):
        ws=wb.add_worksheet(name)
        ws.freeze_panes(1,5); ws.hide_gridlines(2); ws.set_row(0,32)
        widths=[36,18,16,66,17,12,12,10,12,30,42,26,48,12,15,20,18,20,85]
        for col,width in enumerate(widths): ws.set_column(col,col,width,fmt_text)
        for col,h in enumerate(headers): ws.write(0,col,h,fmt_head)
        for ri,r in enumerate(rs,1):
            ws.set_row(ri,82)
            vals=values(r)
            for ci,val in enumerate(vals):
                if val is None or val=='': ws.write_blank(ri,ci,None,fmt_text)
                elif ci==12 and str(val).startswith('https://'): ws.write_url(ri,ci,val,fmt_link,val)
                else: ws.write(ri,ci,val,fmt_money if ci==16 else fmt_num if ci==15 else fmt_text)
            if not r['original_model']:
                ws.write(ri,10,r['selection_status'],fmt_pending)
        ws.autofilter(0,0,len(rs),len(headers)-1)
        ws.set_landscape(); ws.set_paper(8); ws.fit_to_pages(1,0); ws.repeat_rows(0)
        return ws
    table('BOM导入_全量' if full else 'BOM导入_C编号',primary)
    if full:
        table('全部43项总表',rows)
        table('费用及包内件',excluded)
        table('仍需定型的项目',pending)
    ws=wb.add_worksheet('导入说明'); ws.hide_gridlines(2)
    ws.set_column(0,0,27,fmt_text); ws.set_column(1,1,120,fmt_text)
    ws.write_row(0,0,['项目','说明'],fmt_head); ws.freeze_panes(1,0)
    for i,line in enumerate(guide,1):
        ws.set_row(i,48); ws.write(i,0,line[0],fmt_text)
        ws.write(i,1,line[1],fmt_money if isinstance(line[1],(int,float)) else fmt_text)
    wb.close()

full_path=HERE/'桌面助手-全量BOM导入V3.xlsx'
lcsc_path=HERE/'立创BOM导入-全部已知C编号V3.xlsx'
make_book(full_path,materials,True)
make_book(lcsc_path,lcsc,False)
for path,rs in [(HERE/'桌面助手-全量BOM导入V3.csv',materials),(HERE/'立创BOM导入-全部已知C编号V3.csv',lcsc)]:
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f); writer.writerow(headers); writer.writerows(values(r) for r in rs)

data.update(schema_version=3,date=DATE,bom=rows,
    export_status='complete_current_list_exported__product_selection_and_engineering_still_open',
    user_delivery_preference='Prepare complete Excel in current LCSC column order; user imports the new online BOM themselves.',
    current_list_rows=len(rows),full_material_import_rows=len(materials),known_lcsc_code_rows=len(lcsc),
    unfinalized_model_rows=len(pending),
    row_replacements={'W04':'merged_into_W03_FIT0585','W05':'replaced_by_W03_18AWG_kit','W06':'replaced_by_W03_18AWG_kit'},
    previous_bom='hardware-selection/bom-v2/README.md',
    workbook='hardware-selection/bom-v3/桌面助手-全量BOM导入V3.xlsx',
    core_import_workbook='hardware-selection/bom-v3/立创BOM导入-全部已知C编号V3.xlsx')
data['totals']={'scenario':'两路DFR0753备选，1套样机与本次MOQ/备件；含定制/打印/费用预留',
    'reference_cash_cny':total,'web_observed_items_cny':groups['本轮网页价'],
    'historical_screenshot_items_cny':groups['历史用户截图价'],'allowances_cny':groups['预算占位'],
    'known_lcsc_code_rows_cash_cny':lcsc_total,'previous_v2_cash_cny':769.56,
    'change_vs_v2_cny':rounded(Decimal(str(total))-Decimal('769.56'))}
data['core_import_limit']='全部9项已知C编号，包括条件候选及MRO；未上传验证。是否下单仍需核对具体适配条件。'
data['new_research']=[
    {'id':'E04','seller':'汉丽电子科技','model':'PCA9685 16路PWM成品模块','url':'https://item.taobao.com/item.htm?id=967408745108','price_cny':8.9,'stock_label':'有货','delivery_cny_observed':5,'observed':'商品标题、品牌HANLI、数量1、页面价格和有货状态；尺寸/电气图纸未核。'},
    {'id':'W02','model':'FIT0030','url':'https://www.dfrobot.com.cn/goods-366.html','price_cny':6,'stock_label':'有货','observed':'10根/包，母母头、20cm、承受电流<1000mA；使用3包。'},
    {'id':'W03','model':'FIT0585','url':'https://www.dfrobot.com.cn/goods-1770.html','price_cny':10,'stock_label':'有货','observed':'18AWG 0.75mm²，红1m黑1m为1套；镀锡铜/硅胶。厂家22A宣称未视为本机已验证额定。'},
]
(HERE/'bom-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

# Check actual saved Excel cells; imported quantities contain no formulas.
checks=[]
assert len(rows)==43 and len(materials)==37 and len(lcsc)==9
assert len({r['id'] for r in rows})==len(rows)
assert sum(Decimal(str(r['line_cash_cny'])) for r in rows)==Decimal(str(total))
assert set(non_import)=={r['id'] for r in excluded}
for path,rs in [(full_path,materials),(lcsc_path,lcsc)]:
    book=load_workbook(path,data_only=False)
    sheet=book.worksheets[0]
    assert [c.value for c in sheet[1]]==headers
    assert sheet.max_row==len(rs)+1
    for i,r in enumerate(rs,2):
        assert sheet.cell(i,5).value==r['lcsc_code'] or not r['lcsc_code'] and sheet.cell(i,5).value is None
        assert sheet.cell(i,6).value+sheet.cell(i,7).value==r['purchase_quantity']
        assert sheet.cell(i,9).value==r['id']
        assert sheet.cell(i,17).value==r['line_cash_cny']
        assert not any(c.data_type=='f' for c in sheet[i])
    checks.append(f'{path.name}: 表头、{len(rs)}行、型号/C号/数量/金额和无公式值核验通过')
assert not any(r['id'] in non_import for r in materials)
assert all(r['lcsc_code'].startswith('C') for r in lcsc)
assert all(r['original_model']=='' and r['lcsc_code']=='' for r in pending)
validation={'date':DATE,'status':'passed','checks':checks+['37项实物与6项费用/包内件不重叠、无漏行','单机量+备损量=本次采购量','未知型号/封装不补造','新线材合并未重复计费'],
    'all_rows':len(rows),'material_rows':len(materials),'known_c_code_rows':len(lcsc),'model_pending_rows':len(pending),
    'totals':data['totals'],'limitations':['未在网页上传验证字段映射','全部物料选型尚未完成','未改变云BOM或执行采购','金额为静态快照，编辑后需重算']}
(HERE/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
