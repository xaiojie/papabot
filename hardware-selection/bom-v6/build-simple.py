"""V6: move five verified items to Taobao; preserve the exact LCSC import."""
from pathlib import Path
import copy, hashlib, json, shutil
from decimal import Decimal
import xlsxwriter
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
PREV = HERE.parent / 'bom-v5'
DATE = '2026-09-08'
NEW = [
    dict(id='E05', name='麦克风成品板', model='SEN0327 / MSM261S4030H0',
         sku='I2S麦克风模块MSM261S4030H0（单规格）', qty=1, unit='块', price=15,
         shop='DFRobot官方品牌店', url='https://item.taobao.com/item.htm?id=610064842890',
         note='已加购；接3.3V，L/R选择声道；沿用原款',
         evidence='官方淘宝详情图：1.8–3.3V，SCK/WS/LR/SD/GND，直径14mm；主图为已焊排针成品。',
         spec_image='https://img.alicdn.com/imgextra/i4/13560478/O1CN01kiHNxZ1FOyRbCdQ3J_!!13560478.jpg',
         freight=6),
    dict(id='E06', name='功放成品板', model='MAX98357A I2S成品模块',
         sku='MAX98357 焊接', qty=1, unit='块', price=5.19,
         shop='信诚新电子', url='https://item.taobao.com/item.htm?id=866632847020&skuId=5886764007198',
         note='已加购；逻辑5V供电，喇叭焊线接螺丝端子；输出两端均不接地',
         evidence='选中SKU与主图：MAX98357A I2S Mono Amp，Vin 2.5–5.5V，4–8Ω 3W，已焊排针及螺丝端子；选后5.19元。',
         freight=2.4),
    dict(id='P03', name='舵机降压', model='DFR0753',
         sku='DC-DC降压模块6~14V转5V8A（单规格）', qty=1, unit='块', price=55,
         shop='DFRobot官方品牌店', url='https://item.taobao.com/item.htm?id=633486468011',
         note='已加购；只需1块，5V/6A整机负载与温升待测',
         evidence='官方淘宝详情：输入6–14V，输出5V±1%，0–8A；>96%效率条件是12V输入满载。与原官网选型一致。',
         freight=6),
    dict(id='W02', name='母母信号线', model='全铜杜邦线20CM母对母40P',
         sku='全铜杜邦线 母对母（20CM) 2.54mm（1排40P）', qty=1, unit='排', price=8.36,
         shop='zave旗舰店', url='https://detail.tmall.com/item.htm?id=608543275703&skuId=6167752202862',
         note='已加购；共40根，替代原30根；仅用于信号及小电流',
         evidence='实际选择全铜、20CM、母母、2.54mm、1排40P，选后8.36元；非起价2.16元款。',
         freight=0),
    dict(id='W03', name='硅胶电源线', model='2×18AWG红黑硅胶并线',
         sku='硅胶红黑并线 2×18AWG(0.75平方)(1米)', qty=1, unit='份', price=4.29,
         shop='wxrkdz旗舰店', url='https://detail.tmall.com/item.htm?id=900626764964&skuId=5923525883949',
         note='已加购；红黑各1米，需自裁焊接并绝缘',
         evidence='实际SKU为1米双芯并线。详情图：18AWG，每芯150/0.08TS、0.75平方毫米、外径2.3mm，镀锡铜、硅橡胶。未据此认定整机载流能力。',
         spec_image='https://img.alicdn.com/imgextra/i4/2207953209393/O1CN01uTruws2JG3HtKcULv_!!2207953209393.jpg',
         freight=0),
]
assert len(NEW) == 5
for x in NEW:
    x.update(cart_added=True, cart_quantity_verified=1, stock='有货', date=DATE,
             verification='商品页选择后加购成功提示；刷新购物车逐项核对链接、规格、数量及价格')

def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n')

(HERE/'evidence').mkdir(exist_ok=True)
audit = dict(date=DATE, authorization='用户明确回复：找到合适款后加入购物车',
             cart_before_count=5, cart_after_count=10, all_existing_quantities_preserved=True,
             added_items=NEW, added_subtotal_cny=87.84, previous_official_subtotal_cny=128,
             saving_before_freight_cny=40.16, order_placed=False, seller_messages_sent=False,
             freight_note='每页显示运费仅作记录；同店合并运费及最终优惠未进入结算页核验，不重复加总。',
             cart_scope_note='原5项为屏幕、PCA9685、充电板、好盈UBEC、普通电池SKU。MG90S不在当前购物车。',
             battery_note='购物车原14.70元普通电池SKU并非已核实的持续≥10A定制包；保留未改，采购报价仍留空。',
             cart_url='https://cart.taobao.com/')
dump(HERE/'evidence/taobao-cart-audit.json', audit)

# The import file is byte-for-byte identical to the verified V5 deliverable.
for name in ['01-立创商城-BOM导入.xlsx','01-立创商城-BOM导入.csv']:
    shutil.copy2(PREV/name, HERE/name)
shutil.copy2(PREV/'evidence/lcsc-match-audit.json', HERE/'evidence/lcsc-match-audit.json')

oldwb = load_workbook(PREV/'02-淘宝天猫-采购清单.xlsx')
ws = oldwb.active
headers = [c.value for c in ws[1]]
rows = [[c.hyperlink.target if c.hyperlink else c.value for c in row] for row in list(ws)[1:]]
assert len(rows) == 6
rows[0][4:7] = ['touglesy旗舰店','https://detail.tmall.com/item.htm?id=616669602001&skuId=5749013252092','原购物车已核对；2.4寸TFT＋EC11，勿按链接标题买OLED']
rows[1][6] = '沿用已选8个180度款；本轮购物车未见，原链接未留存'
rows[2][6] = '原购物车已有1块；成品板，逻辑接3.3V'
rows[3][6] = '先确认定制；车内14.70元普通款未核实≥10A、均衡和XT30要求'
rows[4][6] = '原购物车已有TYPE-C款；实际充电电流待核'
rows[5][5] = 'https://item.taobao.com/item.htm?id=820639316179&skuId=5525930664079'
rows[5][6] = '原购物车已有1个；到货跳帽设5V，测压后接电子模块'
rows.extend([[x['name'],x['sku'],f"{x['qty']}{x['unit']}",x['price'],x['shop'],x['url'],x['note']] for x in NEW])

name = '02-淘宝天猫-采购清单.xlsx'
wb = xlsxwriter.Workbook(HERE/name, {'strings_to_urls':False})
sheet = wb.add_worksheet('采购清单')
base = {'font_name':'PingFang SC','font_size':11,'valign':'vcenter','text_wrap':True}
normal = wb.add_format(base)
head = wb.add_format({**base,'bold':True,'font_color':'white','bg_color':'#365A50'})
link = wb.add_format({**base,'font_color':'#176B87','underline':1})
money = wb.add_format({**base,'num_format':'0.00'})
sheet.hide_gridlines(2); sheet.freeze_panes(1,0)
sheet.write_row(0,0,headers,head); sheet.set_row(0,30)
for i,width in enumerate([18,43,10,15,22,13,48]): sheet.set_column(i,i,width,normal)
for ri,row in enumerate(rows,1):
    sheet.set_row(ri,50)
    for ci,val in enumerate(row):
        if val is None or val=='': sheet.write_blank(ri,ci,None,normal)
        elif ci==5: sheet.write_url(ri,ci,val,link,'打开商品')
        else: sheet.write(ri,ci,val,money if ci==3 else normal)
sheet.autofilter(0,0,len(rows),6)
sheet.set_landscape(); sheet.set_paper(9); sheet.fit_to_pages(1,1)
sheet.set_margins(.25,.25,.35,.35); sheet.repeat_rows(0)
wb.close()

data = copy.deepcopy(json.loads((PREV/'bom-data.json').read_text()))
by = {r['id']:r for r in data['bom']}
old_totals = copy.deepcopy(data['totals'])
for x in NEW:
    r = by[x['id']]
    r['previous_official_selection'] = copy.deepcopy(r)
    r.update(model=x['model'],original_model=x['model'],channel='淘宝/天猫',url=x['url'],
             original_brand=x['shop'],purchase_quantity=1,installed_quantity=1,unit=x['unit'],
             unit_price_or_allowance_cny=x['price'],line_cash_cny=x['price'],
             price_basis='2026-09-08实际SKU选择价与购物车价一致，不含运费',
             procurement_file=name,selection_status='已选并加购，购物车数量1已核验',
             selected_for_prototype=True,hardware_tested=False,purchase_confirmed=False,
             cart_added=True,cart_quantity_verified=1,exact_purchase_sku=x['sku'],
             source_evidence='evidence/taobao-cart-audit.json')
    if x['id']=='E06':
        r['notes']='替代DFR0954，选MAX98357焊接款。板上MAX98357A、2.5–5.5V、I2S，接逻辑5V。喇叭1.25mm插头不能直插螺丝端子，应改焊线并绝缘；BTL输出两端都不接地。'
        r['mechanical_constraints']='新板框尺寸、孔距、排针高度及螺丝端子空间待实测；不得套用DFR0954的接口/尺寸，CAD未修改。'
    elif x['id']=='W02':
        r['item']='全铜母对母20cm信号杜邦线，1排40根'
        r['notes']='1排40P满足原30根需求；多出的10根备用。2.54mm全铜母母20cm，仅信号及小电流；不作舵机/电池/逻辑总电源线。不沿用原FIT0030电流参数。'
        r['included_leads']=40; r['required_leads']=30; r['spare_leads']=10
    elif x['id']=='W03':
        r['item']='红黑2×18AWG硅胶并线，1米（每色1米）'
        r['notes']='每芯镀锡铜0.75平方毫米，18AWG，红黑双芯1米。需自裁焊接绝缘；复用舵机原线，XT30仍要求随电池提供；线长、压降和成束温升待验证。'
    else:
        r['notes'] += ' 本轮通过DFRobot官方淘宝店加购原款；未下单。'
by['E01']['url']=rows[0][5]
by['E01']['cart_quantity_verified']=1
by['P04']['url']=rows[5][5]
by['P01']['cart_verification_note']=audit['battery_note']
by['P01']['cart_verified_as_custom_spec']=False

part=data['procurement_partition']
part['taobao_tmall'] += part['other_official_sites']
part['other_official_sites']=[]
assert len(part['taobao_tmall'])==11
data.update(schema_version=6,previous_bom='hardware-selection/bom-v5/README.md',
            workbook='hardware-selection/bom-v6/01-立创商城-BOM导入.xlsx',
            core_import_workbook='hardware-selection/bom-v6/01-立创商城-BOM导入.xlsx',
            export_status='two_simple_single_sheet_files__all_five_official_rows_moved_to_taobao',
            source_match_audit='hardware-selection/bom-v6/evidence/lcsc-match-audit.json',
            cart_audit='hardware-selection/bom-v6/evidence/taobao-cart-audit.json')
data['procurement_files']=[dict(path='hardware-selection/bom-v6/01-立创商城-BOM导入.xlsx',rows=10,columns=3,sheets=1),dict(path='hardware-selection/bom-v6/'+name,rows=11,columns=7,sheets=1)]
data['procurement_totals'].update(taobao_known_prices_cny=258.27,other_official_sites_cny=0,added_five_subtotal_cny=87.84,
                                five_saving_before_freight_cny=40.16,known_selected_purchase_cny=369.64)
data['historical_v5_totals']=old_totals
data['totals'].update(reference_cash_cny=721.02,web_observed_items_cny=248.49,
                      closeout_items_without_retained_servo_regulator_cny=85.97,
                      closeout_items_with_retained_servo_regulator_cny=140.97,
                      previous_v5_cash_cny=761.18,change_vs_v5_cny=-40.16,
                      note='工程预算延续旧立创价格和346元占位；实际采购已知价格见procurement_totals，不含未报价电池及运费。')
data['procurement_totals']['engineering_budget_snapshot_cny']=721.02
data['selection_closeout']['taobao_five_added']=True
dump(HERE/'bom-data.json',data)

checks=[]
for path,expected_rows,expected_cols in [(HERE/'01-立创商城-BOM导入.xlsx',10,3),(HERE/name,11,7)]:
    check=load_workbook(path); sh=check.active
    assert len(check.worksheets)==1 and sh.max_row==expected_rows+1 and sh.max_column==expected_cols
    assert not sh.merged_cells.ranges
    assert not any(c.data_type=='f' for row in sh for c in row)
    if expected_cols==7:
        assert [c.value for c in sh[1]]==headers
        for ri,row in enumerate(rows,2):
            for ci,value in enumerate(row,1):
                cell=sh.cell(ri,ci)
                if ci==6 and value: assert cell.hyperlink.target==value
                else: assert cell.value==(value or None)
    checks.append(f'{path.name}: 单页、{expected_rows}项、{expected_cols}列，通过')
assert (HERE/'01-立创商城-BOM导入.xlsx').read_bytes()==(PREV/'01-立创商城-BOM导入.xlsx').read_bytes()
assert sum(Decimal(str(x['price'])) for x in NEW)==Decimal('87.84')
assert len(data['bom'])==len(set(sum(part.values(),[])))==42
validation=dict(status='passed',date=DATE,checks=checks,lcsc_import_byte_identical_to_v5=True,
                lcsc_sha256=hashlib.sha256((HERE/'01-立创商城-BOM导入.xlsx').read_bytes()).hexdigest(),
                coverage=dict(lcsc=10,taobao=11,official=0,deferred_or_nonpurchase=21,total=42),
                newly_added_to_taobao=5,cart_quantities_verified=True,cart_count=10,
                order_placed=False,online_lcsc_bom_modified=False)
dump(HERE/'validation.json',validation)
print(json.dumps(validation,ensure_ascii=False,indent=2))
