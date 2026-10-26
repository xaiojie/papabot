"""Build three simple, single-sheet purchasing files from verified sourcing data.

Only C codes verified against the user's actual LCSC matching export are allowed
into the LCSC import file. No online BOM changes or uploads are performed.
"""
from pathlib import Path
from decimal import Decimal
import copy
import json
import csv
import xlsxwriter
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
data = copy.deepcopy(json.loads((ROOT/'hardware-selection/bom-v4/bom-data.json').read_text()))
audit = json.loads((HERE/'evidence/lcsc-match-audit.json').read_text())
by = {r['id']: r for r in data['bom']}
lcsc_ids = [r['id'] for r in audit['retained_matching_rows']]
taobao_ids = ['E01','E02','E04','P01','P02','P04']
official_ids = ['E05','E06','P03','W02','W03']
listed_ids = lcsc_ids+taobao_ids+official_ids
assert len(listed_ids) == len(set(listed_ids))
deferred_ids = [r['id'] for r in data['bom'] if r['id'] not in listed_ids]

lcsc_rows = []
for match in audit['retained_matching_rows']:
    r = by[match['id']]
    assert match['original_code'] == match['matched_code'] == r['lcsc_code']
    assert match['match_status'] == '系统匹配' and match['channel'] == '现货'
    qty = int(match['purchase_quantity'])
    assert qty == r['purchase_quantity'] and qty >= int(match['moq'])
    assert int(match['stock']) >= qty
    # Correct W01's old brand-prefixed pseudo-model to the actual manufacturer model.
    r['model'] = r['original_model'] = match['matched_model']
    r['original_brand'] = match['matched_brand']
    r['procurement_file'] = '01-立创商城-BOM导入.xlsx'
    r['lcsc_import_match_verified'] = True
    r['lcsc_import_match_evidence'] = 'evidence/lcsc-match-audit.json'
    lcsc_rows.append([match['matched_code'], match['matched_model'], qty])

taobao_details = {
    'E01': ['TFT屏幕＋旋钮','TOUGLESY 2.4英寸 ST7789＋EC11','1块',32.4,'原已选店铺','','沿用已选购物车款；原链接未留存'],
    'E02': ['舵机','MG90S 金属齿轮180度定位款','8个',10,'原已选店铺','','沿用已选款，含舵盘螺丝；原链接未留存'],
    'E04': ['PWM成品板','HANLI PCA9685 16路成品模块','1块',8.9,'汉丽电子科技',by['E04']['url'],'购买成品板；逻辑接3.3V'],
    'P01': ['定制电池包','2S 7.4V 1000mAh，持续≥10A','1包',None,'明盛电子（定制入口）',by['P01']['url'],'先询价；含每节保护、均衡及XT30配对件'],
    'P02': ['充电模块','MK021163，选择TYPE-C接口','1块',14.13,'TELESKY旗舰店',by['P02']['url'],'2S/8.4V款；实际充电电流需核对'],
    'P04': ['逻辑降压','好盈86010010，选择3A(2-6S)','1个',35,'顶峰模型',by['P04']['url'],'到货将跳帽设为5V，测压后接电子模块'],
}
official_details = {
    'E05': ['麦克风成品板','SEN0327 I²S麦克风模块','1块',15,'DFRobot官网',by['E05']['url'],'3.3V供电；不要买INMP441裸芯片'],
    'E06': ['功放成品板','DFR0954 / MAX98357A','1块',30,'DFRobot官网',by['E06']['url'],'喇叭1.25mm插头不直配PH2.0，按焊线连接'],
    'P03': ['舵机降压','DFR0753，6–14V转5V','1块',55,'DFRobot官网',by['P03']['url'],'只买1块；整机负载与温升待测'],
    'W02': ['母母信号线','FIT0030，20cm，10根/包','3包',6,'DFRobot官网',by['W02']['url'],'共30根，仅信号及小电流；不作总电源线'],
    'W03': ['硅胶电源线','FIT0585，18AWG红黑各1米','1套',10,'DFRobot官网',by['W03']['url'],'需自行裁切、焊接和绝缘'],
}

imports = ['原C编码','原型号','单套用量']
external = ['配件','购买规格','数量','参考单价(元)','店铺/官网','商品链接','下单备注']
books = [
    ('01-立创商城-BOM导入.xlsx', imports, lcsc_rows, [19,34,14]),
    ('02-淘宝天猫-采购清单.xlsx', external, [taobao_details[i] for i in taobao_ids], [18,35,10,15,21,13,44]),
    ('03-产品官网-采购清单.xlsx', external, [official_details[i] for i in official_ids], [18,35,10,15,21,13,44]),
]

def write_book(name, headers, rows, widths):
    wb = xlsxwriter.Workbook(HERE/name, {'strings_to_urls':False})
    base = {'font_name':'PingFang SC','font_size':11,'valign':'vcenter','text_wrap':True}
    normal = wb.add_format(base)
    number = wb.add_format({**base,'num_format':'0.00'})
    head = wb.add_format({**base,'bold':True,'font_color':'#FFFFFF','bg_color':'#365A50'})
    link = wb.add_format({**base,'font_color':'#176B87','underline':1})
    ws = wb.add_worksheet('采购清单')
    ws.hide_gridlines(2); ws.freeze_panes(1,0)
    ws.set_row(0,30); ws.write_row(0,0,headers,head)
    for ci,width in enumerate(widths): ws.set_column(ci,ci,width,normal)
    for ri,row in enumerate(rows,1):
        ws.set_row(ri,32 if len(headers)==3 else 48)
        for ci,value in enumerate(row):
            if value is None or value=='': ws.write_blank(ri,ci,None,normal)
            elif ci==5: ws.write_url(ri,ci,value,link,'打开商品')
            else: ws.write(ri,ci,value,number if ci==3 else normal)
    ws.autofilter(0,0,len(rows),len(headers)-1)
    ws.set_landscape(); ws.set_paper(9); ws.fit_to_pages(1,1)
    ws.set_margins(.25,.25,.35,.35); ws.repeat_rows(0)
    wb.close()

checks = []
for name,headers,rows,widths in books:
    write_book(name,headers,rows,widths)
    wb = load_workbook(HERE/name,data_only=False)
    assert len(wb.worksheets)==1
    ws = wb.worksheets[0]
    assert ws.max_row==len(rows)+1 and ws.max_column==len(headers)
    assert not ws.merged_cells.ranges
    assert [c.value for c in ws[1]]==headers
    for ri,row in enumerate(rows,2):
        for ci,val in enumerate(row,1):
            cell = ws.cell(ri,ci)
            assert cell.data_type!='f'
            if len(headers)>3 and ci==6 and val:
                assert cell.hyperlink.target==val and cell.value=='打开商品'
            else:
                assert cell.value==(None if val=='' else val),(name,ri,ci)
    checks.append(f'{name}: 单页、{len(rows)}项、{len(headers)}列、数值及链接检查通过')
with (HERE/'01-立创商城-BOM导入.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.writer(f); writer.writerow(imports); writer.writerows(lcsc_rows)

for rid in taobao_ids: by[rid]['procurement_file']='02-淘宝天猫-采购清单.xlsx'
for rid in official_ids: by[rid]['procurement_file']='03-产品官网-采购清单.xlsx'
for rid in deferred_ids: by[rid]['procurement_file']=None
data.update(schema_version=5,previous_bom='hardware-selection/bom-v4/README.md',
    workbook='hardware-selection/bom-v5/01-立创商城-BOM导入.xlsx',
    core_import_workbook='hardware-selection/bom-v5/01-立创商城-BOM导入.xlsx',
    export_status='simple_single_sheet_files_separated_by_purchase_channel',
    full_material_import_rows=10, known_lcsc_code_rows=10,
    user_delivery_preference='All Excel files stay simple. LCSC import contains only verified matching stocked C codes. Taobao/Tmall and manufacturer websites have separate one-sheet lists. User imports.',
    core_import_limit='10项均在用户当前V4云BOM导出结果中匹配到相同C编号且为现货；本次简洁文件尚未由助手上传。')
data['procurement_partition'] = dict(lcsc=lcsc_ids,taobao_tmall=taobao_ids,
    other_official_sites=official_ids,deferred_or_nonpurchase=deferred_ids)
data['procurement_files'] = [dict(path='hardware-selection/bom-v5/'+name,rows=len(rs),columns=len(h),sheets=1) for name,h,rs,w in books]
data['procurement_totals'] = dict(lcsc_current_account_quote_cny=audit['account_quote_subtotal_cny'],
    taobao_known_prices_cny=170.43,taobao_battery_quote_cny=None,other_official_sites_cny=128,
    exclude='battery unquoted amount, other unselected materials, PCB/printing, freight and contingency',
    engineering_budget_snapshot_cny=data['totals']['reference_cash_cny'])
data['source_match_audit']='hardware-selection/bom-v5/evidence/lcsc-match-audit.json'
assert len(data['bom'])==len(set(listed_ids+deferred_ids))==42
assert len(lcsc_rows)==10 and all(row[0].startswith('C') and row[2]>0 for row in lcsc_rows)
assert not any(r in lcsc_ids for r in ['E05','E06','P03','W02','W07','W10','P04'])
(HERE/'bom-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
validation=dict(status='passed',date='2026-09-08',checks=checks,
    lcsc_verified_exact_code_matches=10,lcsc_stocked_rows=10,lcsc_columns=3,
    lcsc_quantity_basis='actual purchase quantities including MOQ/spares; import with procurement sets=1, no additional spares',
    coverage=dict(lcsc=10,taobao=6,official_sites=5,deferred_or_nonpurchase=21,total=42),
    new_online_import_performed=False,current_online_bom_modified=False,order_placed=False,
    source_match_audit='evidence/lcsc-match-audit.json')
(HERE/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
