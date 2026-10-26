"""Export the current local sourcing data; never uploads or changes online BOMs.

Edit bom-data.json then run this file. select-v4.py recreates the dated baseline
from V3 and must only be used when deliberately resetting V4 sourcing data.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv
import json
import math
import unicodedata
import xlsxwriter
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
data = json.loads((HERE / 'bom-data.json').read_text())
rows = data['bom']

def money(value):
    return float(Decimal(str(value)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))

def cash(rs):
    return money(sum(Decimal(str(r['line_cash_cny'])) for r in rs))

for r in rows:
    r['spare_quantity'] = max(0, r['purchase_quantity'] - r['installed_quantity'])
    if r['export_to_full_material_sheet']:
        assert r['purchase_quantity'] >= r['installed_quantity']
    r['line_cash_cny'] = money(Decimal(str(r['purchase_quantity'])) * Decimal(str(r['unit_price_or_allowance_cny'])))

materials = [r for r in rows if r['export_to_full_material_sheet']]
excluded = [r for r in rows if not r['export_to_full_material_sheet']]
lcsc = [r for r in materials if r['lcsc_code']]
pending = [r for r in materials if not r['original_model']]
selected = [r for r in materials if r['id'] in data['selection_closeout']['row_ids']]
groups = {basis: cash([r for r in rows if r['price_basis'] == basis])
          for basis in ['本轮网页价', '历史用户截图价', '预算占位', '已含在其他行']}
total = cash(rows)
data.update(current_list_rows=len(rows), full_material_import_rows=len(materials),
            known_lcsc_code_rows=len(lcsc), unfinalized_model_rows=len(pending))
data['totals'] = dict(
    scenario='舵机1块DFR0753＋逻辑1块好盈86010010＋S1A机械总开关；1套样机含MOQ/备件/试制预留',
    reference_cash_cny=total, web_observed_items_cny=groups['本轮网页价'],
    historical_screenshot_items_cny=groups['历史用户截图价'], allowances_cny=groups['预算占位'],
    known_lcsc_code_rows_cash_cny=cash(lcsc), previous_v3_cash_cny=761.76,
    change_vs_v3_cny=money(Decimal(str(total)) - Decimal('761.76')),
    closeout_items_without_retained_servo_regulator_cny=cash([r for r in selected if r['id'] != 'P03']),
    closeout_items_with_retained_servo_regulator_cny=cash(selected))

headers = ['原型号','原品牌','原封装','原参数','原C编码','单套用量','备损量',
    '单位','物料编号','配件名称','选型状态','采购渠道','商品链接','已查MOQ','建议采购数量',
    '参考单价/预留(元)','参考现金(元)','价格依据','备注']

def values(r):
    params = f"{r['item']}；{r['model']}；单位：{r['unit']}"
    if r['id'] == 'P04':
        params += '；购买3A(2-6S)，到货跳帽置5V档'
    if not r['original_model']:
        params = '待定规格，仅列需求：' + params
    return [r['original_model'],r['original_brand'],r['original_package'],params,
        r['lcsc_code'],r['installed_quantity'],r['spare_quantity'],r['unit'],r['id'],
        r['item'],r['selection_status'],r['channel'],r['url'],r['observed_moq'],
        r['purchase_quantity'],r['unit_price_or_allowance_cny'],r['line_cash_cny'],
        r['price_basis'],r['notes']+' '+r['mechanical_constraints']]

guide = [
    ['日期',data['date']],
    ['本轮结论','PCA9685成品、低价逻辑降压、直流总开关、小包装信号及电源线已全部选定为样机采购方案；舵机路仅保留1块DFR0753。'],
    ['范围',f'{len(rows)}项总表；首表{len(materials)}项实物需求；{len(lcsc)}项有C编号；另{len(pending)}项为此前电池定制/其他辅料/电路结构依赖，不是本轮四类商品未选。'],
    ['快速采购','查看“本轮定型采购”：含5行新定商品＋1行保留舵机电源；所有SKU/包装、数量、网址与接线条件已写入。'],
    ['导入步骤1','用户在新版立创BOM导入Excel，选择首表“BOM导入_全量”，第1行作为表头。'],
    ['导入步骤2','按前7列依次映射原型号、原品牌、原封装、原参数、原C编码、单套用量、备损量；采购套数设1。'],
    ['数量口径','单套用量＋备损量＝建议采购数量；备损量是绝对数量，含备件及MOQ多买余量。不要把建议采购数量重复叠加。'],
    ['数量示例','PM2X6：20＋80＝100个；1000µF：2＋3＝5个。FIT0030用量3表示3包，共30根；FIT0585用量1表示1套红黑各1米。'],
    ['若备损按百分比','若网页备损量是百分比，不映射本列；仅就1套采购，将建议采购数量映射到单套用量并将备损置0。增加套数须重算MOQ/备件。'],
    ['C编号',f'可单用“立创BOM导入-全部已知C编号V4.xlsx”，{len(lcsc)}项，参考{cash(lcsc):.2f}元；本轮新增S1A / C5366606。含条件候选和MRO，自动匹配及适配需核对。'],
    ['外采匹配','淘宝/DFRobot外采不编造C编号；选中好盈3A(2-6S)，不是同链接5A/10A/25A变体。不接受系统推荐的裸芯片替代成品模块。'],
    ['封装与空型号','封装只填已核字段；未定项目原型号留空，需求写入原参数。导入预览须核对空型号行是否保留，完整总表始终含全部记录。'],
    ['两路电源','DFR0753×1给舵机；好盈86010010×1给逻辑/音频。输入共用受控2S母线，共地；两个5V输出不得并联。'],
    ['5V跳帽','好盈86010010支持5V/6V切换，必须断电设5V并测量后接电子模块；逻辑路设计目标5V/2A，厂家标称3A，不表示本机温升/纹波已测。'],
    ['总开关','S1A保持型直接断开两路工作负载，串在保险后电池正极；充电支路接在开关前的受保护节点。首版关机充电，USB调试供电也须单独考虑回灌。'],
    ['低电量','移除原F05 MOSFET控制预算；S1A不自动断电。暂建议约6.4V负载电压停止动作并提示手动关机，阈值/回差未验证，每节电池保护仍由合格电池包承担。'],
    ['线束边界','FIT0030是成品信号线；FIT0585是需自裁焊接的硅胶线。复用MG90S原线和电池包含的XT30配对件。载板排针/热缩管仍在原有辅料需求行，不表示预压接整机线束已加工完成。'],
    ['实物检查','PCA VCC/I²C上拉均3.3V，默认OE禁用；舵机电流外部分配。核实低电量压降、温升、噪声、模块尺寸和线束布局，再装入外壳。'],
    ['完整参考现金',total],
    ['网页观察价',groups['本轮网页价']],
    ['历史截图价',groups['历史用户截图价']],
    ['预算占位',groups['预算占位']],
    ['价格范围','观察价为2026-09-08快照；定制/PCB/打印/其他辅料仍含预留。运费统一留在L01，不重复加各店运费。不是整单结算报价。'],
    ['本轮金额变化','逻辑降压省20元；总开关从5元预留改29.42元；删除F05的5元预留，净减少0.58元。'],
    ['文件与重建','导入单元格均为静态值。改bom-data.json后运行build-v4.py重算；select-v4.py仅用于从V3重建本轮基线，会重置V4手工改动。'],
]

def make_book(path, primary, full):
    wb = xlsxwriter.Workbook(path, {'strings_to_urls': False})
    base = {'font_name':'PingFang SC','font_size':10,'text_wrap':True,'valign':'top'}
    text_fmt = wb.add_format(base)
    head_fmt = wb.add_format({**base,'bold':True,'font_color':'#FFFFFF','bg_color':'#365A50','valign':'vcenter'})
    money_fmt = wb.add_format({**base,'num_format':'0.00'})
    price_fmt = wb.add_format({**base,'num_format':'0.######'})
    link_fmt = wb.add_format({**base,'font_color':'#176B87','underline':1})
    pending_fmt = wb.add_format({**base,'bg_color':'#FFF2CC','font_color':'#7F6000'})
    chosen_fmt = wb.add_format({**base,'bg_color':'#E2EFDA','font_color':'#285638'})
    def table(name, rs):
        ws = wb.add_worksheet(name)
        ws.freeze_panes(1,5); ws.hide_gridlines(2); ws.set_row(0,32)
        widths = [36,23,14,64,17,12,12,10,12,32,38,25,48,12,15,20,18,20,90]
        for ci,width in enumerate(widths): ws.set_column(ci,ci,width,text_fmt)
        ws.write_row(0,0,headers,head_fmt)
        for ri,r in enumerate(rs,1):
            row_values = values(r)
            wrapped_lines = []
            for ci,value in enumerate(row_values):
                width_units = sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in str(value or ''))
                wrapped_lines.append(max(1, math.ceil(width_units / max(8,widths[ci]-5))))
            ws.set_row(ri,min(360,max(82,max(wrapped_lines)*14+12)))
            for ci,value in enumerate(row_values):
                fmt = money_fmt if ci == 16 else price_fmt if ci == 15 else text_fmt
                if ci == 10:
                    fmt = chosen_fmt if r.get('selected_for_prototype') else pending_fmt if not r['original_model'] else text_fmt
                if value is None or value == '': ws.write_blank(ri,ci,None,fmt)
                elif ci == 12 and str(value).startswith('https://'): ws.write_url(ri,ci,value,link_fmt,value)
                else: ws.write(ri,ci,value,fmt)
        ws.autofilter(0,0,len(rs),len(headers)-1)
        ws.set_landscape(); ws.set_paper(8); ws.fit_to_pages(1,0); ws.repeat_rows(0)
    table('BOM导入_全量' if full else 'BOM导入_C编号', primary)
    if full:
        table('本轮定型采购', selected)
        table(f'全部{len(rows)}项总表', rows)
        table('费用及包内件', excluded)
        table('其他待定需求', pending)
    ws = wb.add_worksheet('导入说明')
    ws.hide_gridlines(2); ws.freeze_panes(1,0)
    ws.set_column(0,0,26,text_fmt); ws.set_column(1,1,122,text_fmt)
    ws.write_row(0,0,['项目','说明'],head_fmt)
    for ri,line in enumerate(guide,1):
        ws.set_row(ri,52); ws.write(ri,0,line[0],text_fmt)
        ws.write(ri,1,line[1],money_fmt if isinstance(line[1],(int,float)) else text_fmt)
    wb.close()

full_path = HERE / '桌面助手-全量BOM导入V4.xlsx'
core_path = HERE / '立创BOM导入-全部已知C编号V4.xlsx'
for path,rs,full in [(full_path,materials,True),(core_path,lcsc,False)]:
    make_book(path,rs,full)
    with path.with_suffix('.csv').open('w',encoding='utf-8-sig',newline='') as file:
        writer = csv.writer(file); writer.writerow(headers); writer.writerows(values(r) for r in rs)
(HERE / 'bom-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

# Reopen actual output files: quantities, SKU selection, complete coverage and cash.
assert len(rows) == len({r['id'] for r in rows})
assert {r['id'] for r in excluded} == {'W08','P09','M11','M12','L01','L02'}
assert len(selected) == 6 and all(r['original_model'] and r.get('selected_for_prototype') for r in selected)
assert sum(r['purchase_quantity'] for r in rows if r['model'] == 'DFR0753') == 1
assert not any(r['id'] == 'F05' for r in rows)
assert [r for r in rows if r['id'] == 'P04'][0]['original_model'] == '86010010'
assert [r for r in rows if r['id'] == 'P06'][0]['lcsc_code'] == 'C5366606'
assert all(r['purchase_quantity'] >= r['observed_moq'] for r in selected)
assert all(r['lcsc_code'].startswith('C') for r in lcsc)
assert all(not r['lcsc_code'] for r in pending)
checks = []
for path,rs in [(full_path,materials),(core_path,lcsc)]:
    book = load_workbook(path,data_only=False)
    sheet = book.worksheets[0]
    assert [c.value for c in sheet[1]] == headers
    assert sheet.max_row == len(rs)+1
    for ri,r in enumerate(rs,2):
        expected = values(r)
        for ci,value in enumerate(expected,1):
            assert sheet.cell(ri,ci).value == (None if value == '' else value), (r['id'],ci)
        assert sheet.cell(ri,6).value + sheet.cell(ri,7).value == r['purchase_quantity']
        assert not any(c.data_type == 'f' for c in sheet[ri])
    if path == full_path:
        all_sheet = book[f'全部{len(rows)}项总表']
        assert all_sheet.max_row == len(rows)+1
        assert sum(Decimal(str(all_sheet.cell(ri,17).value)) for ri in range(2,len(rows)+2)) == Decimal(str(total))
        assert book['本轮定型采购'].max_row == len(selected)+1
    checks.append(f'{path.name}: {len(rs)}行，19列表头、全部单元格、数量和静态金额核验通过')
validation = dict(date=data['date'],status='passed',all_rows=len(rows),material_rows=len(materials),
    known_c_code_rows=len(lcsc),other_model_pending_rows=len(pending),selected_categories=4,
    checks=checks+['仅1块DFR0753，逻辑86010010，S1A C码正确，F05不重复计价',
                   '本轮6行采购数量均满足已查MOQ；30根信号线按3包计量',
                   '6项费用及包内件未混入实物自动匹配表，全量成本未漏项'],
    totals=data['totals'],limitations=['未上传云BOM验证','未购买或改购物车','整机电气/机械实测尚未执行','其他电池定制及电路结构依赖仍保留'])
(HERE / 'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
